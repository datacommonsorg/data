# Copyright 2026 Google LLC
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#      http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.
"""Unit tests for AustraliaStatistics preprocessing and import scripts."""

import datetime
import io
import tempfile
import unittest

import openpyxl

from preprocess_abs_labour_force import CSV_COLUMNS
from preprocess_abs_labour_force import SERIES_COLUMNS
from preprocess_abs_labour_force import format_obs_date
from preprocess_abs_labour_force import format_series_value
from preprocess_abs_labour_force import merge_employment_records
from preprocess_abs_labour_force import parse_labour_force_workbook
from preprocess_abs_life_expectancy import build_do001_periods
from preprocess_abs_life_expectancy import ensure_pvmap_periods
from preprocess_abs_life_expectancy import parse_do001_national_age_table
from preprocess_abs_life_expectancy import parse_do002_states_table


def _build_sample_labour_force_workbook() -> bytes:
    """Builds an in-memory openpyxl workbook matching ABS 6202.0 Data1 layout."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Data1'

    titles = [''] + SERIES_COLUMNS
    units = [
        '',
        '000',
        '000',
        '000',
        '000',
        '000',
        '000',
        'Percent',
        'Percent',
        'Percent',
    ]
    series_types = [''] + ['Original'] * len(SERIES_COLUMNS)

    ws.append(titles)
    ws.append(units)
    ws.append(series_types)
    for _ in range(7):
        ws.append(['Metadata'] + [''] * len(SERIES_COLUMNS))

    ws.append([
        datetime.datetime(2026, 2, 1),
        14500.25,
        7500.1,
        7000.15,
        600.5,
        320.25,
        280.25,
        4.14,
        4.26,
        4.01,
    ])
    ws.append([
        datetime.datetime(2026, 3, 1),
        14550.0,
        7520.0,
        7030.0,
        610.0,
        325.0,
        285.0,
        4.19,
        4.32,
        4.05,
    ])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _build_sample_do001_workbook() -> bytes:
    """Builds an in-memory workbook matching ABS 3302055001DO001 Table 9."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Table 9'
    for _ in range(7):
        ws.append(['Header'])
    ws.append(
        [0, 100000, 0.003, 99700, 81.06156, 100000, 0.003, 99727, 85.1047])
    ws.append(
        [1, 99647, 0.0002, 99634, 80.34817, 99696, 0.0002, 99684, 84.3641])
    ws.append([99, 2500, 0.31, 2000, 2.381, 5500, 0.29, 4500, 2.619])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _build_sample_do002_workbook() -> bytes:
    """Builds an in-memory workbook matching ABS 3302055001DO002 Table 1."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Table 1'
    for _ in range(5):
        ws.append([None])
    ws.append([None, None, '2021–2023', None, None, None, '2022–2024', None])
    ws.append([None, None, 'Life expectancy at birth'])
    ws.append(
        [None, None, 'Males', 'Females', 'Persons', None, 'Males', 'Females'])
    ws.append([None, None, 'years', 'years', 'years', None, 'years', 'years'])
    ws.append(['ASGS', 'GCCSA, SA4'])
    ws.append(
        [None, 'Total New South Wales', 81.2, 85.2, 83.1, None, 81.2, 85.1])
    ws.append([None, 'Australia', 81.1, 85.1, 83.0, None, 81.1, 85.1])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class AustraliaImportTest(unittest.TestCase):
    """Tests ABS workbook parsing and historical baseline merging."""

    def test_format_series_value(self):
        self.assertEqual(format_series_value(13756.3407, '000'), '13756340.7')
        self.assertEqual(format_series_value(3.74, 'Percent'), '3.7')

    def test_format_obs_date(self):
        self.assertEqual(format_obs_date(datetime.datetime(2026, 3, 1)),
                         '2026-03')
        self.assertEqual(format_obs_date('2025-11-01'), '2025-11')
        self.assertIsNone(format_obs_date('Series ID'))

    def test_parse_labour_force_workbook(self):
        wb_bytes = _build_sample_labour_force_workbook()
        rows = parse_labour_force_workbook(wb_bytes, 'country/AUS')
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[0]['observationDate'], '2026-02')
        self.assertEqual(rows[0]['observationAbout'], 'country/AUS')
        self.assertEqual(rows[0]['Employed total ;  Persons ;'], '14500250')
        self.assertEqual(rows[0]['Unemployment rate ;  Persons ;'], '4.1')
        self.assertEqual(rows[1]['observationDate'], '2026-03')
        self.assertEqual(rows[1]['Employed total ;  Persons ;'], '14550000')

    def test_merge_employment_records_preserves_baseline(self):
        baseline_row = {col: '100' for col in CSV_COLUMNS}
        baseline_row['observationDate'] = '2023-01'
        baseline_row['observationAbout'] = 'country/AUS'

        revised_old_row = dict(baseline_row)
        revised_old_row['Employed total ;  Persons ;'] = '999999'

        new_row = dict(baseline_row)
        new_row['observationDate'] = '2026-03'
        new_row['Employed total ;  Persons ;'] = '14550000'

        merged = merge_employment_records([baseline_row],
                                          [revised_old_row, new_row])
        self.assertEqual(len(merged), 2)
        self.assertEqual(merged[0]['observationDate'], '2023-01')
        self.assertEqual(merged[0]['Employed total ;  Persons ;'], '100')
        self.assertEqual(merged[1]['observationDate'], '2026-03')
        self.assertEqual(merged[1]['Employed total ;  Persons ;'], '14550000')

    def test_parse_life_expectancy_workbooks(self):
        age_map = parse_do001_national_age_table(_build_sample_do001_workbook())
        self.assertNotIn(('Male', '0'), age_map)
        self.assertEqual(age_map[('Male', '1')], '80.35')
        self.assertEqual(age_map[('Female', '1')], '84.36')
        self.assertEqual(age_map[('Female', '99')], '2.62')

        periods, state_vals = parse_do002_states_table(
            _build_sample_do002_workbook())
        self.assertEqual(periods, ['2021-2023', '2022-2024'])
        self.assertEqual(state_vals[('Male', 'wikidataId/Q3224', '2022-2024')],
                         '81.2')
        self.assertEqual(state_vals[('Female', 'country/AUS', '2022-2024')],
                         '85.1')

    def test_dynamic_periods_and_pvmap_update(self):
        periods = build_do001_periods('2023-2025', start_year=2021)
        self.assertEqual(periods, ['2021-2023', '2022-2024', '2023-2025'])
        with tempfile.NamedTemporaryFile('w+', suffix='.py', delete=True) as tf:
            tf.write('{\n   "2021-2023": {"observationDate": "2023"},\n}\n')
            tf.flush()
            ensure_pvmap_periods(tf.name, periods)
            tf.seek(0)
            updated = tf.read()
            self.assertIn('"2022-2024"', updated)
            self.assertIn('"2023-2025"', updated)
            self.assertIn('"observationDate": "2025"', updated)


if __name__ == '__main__':
    unittest.main()
