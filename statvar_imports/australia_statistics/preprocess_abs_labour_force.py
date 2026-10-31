# Copyright 2026 Google LLC
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#      http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.
"""Downloads and merges ABS Labour Force workbooks into aus_employment_int.csv."""

import calendar
import csv
import datetime
import io
import os
import re
from typing import Dict, List, Optional, Tuple

from absl import app
from absl import flags
from absl import logging
from bs4 import BeautifulSoup
import openpyxl
import requests
from requests.adapters import HTTPAdapter
from urllib3.util.retry import Retry

FLAGS = flags.FLAGS

flags.DEFINE_string(
    'employment_csv',
    os.path.join(
        os.path.dirname(os.path.abspath(__file__)),
        'intermediate_files',
        'aus_employment_int.csv',
    ),
    'Path to intermediate aus_employment_int.csv to update.',
)
flags.DEFINE_string(
    'release_url',
    '',
    'Optional explicit ABS Labour Force release URL.',
)

LATEST_RELEASE_URL = (
    'https://www.abs.gov.au/statistics/labour/employment-and-unemployment/'
    'labour-force-australia/latest-release')
BASE_RELEASE_PREFIX = (
    'https://www.abs.gov.au/statistics/labour/employment-and-unemployment/'
    'labour-force-australia')

WORKBOOK_PLACES: List[Tuple[str, str]] = [
    ('6202001.xlsx', 'country/AUS'),
    ('6202004.xlsx', 'wikidataId/Q3224'),
    ('6202005.xlsx', 'wikidataId/Q36687'),
    ('6202006.xlsx', 'wikidataId/Q36074'),
    ('6202007.xlsx', 'wikidataId/Q35715'),
    ('6202008.xlsx', 'wikidataId/Q3206'),
    ('6202009.xlsx', 'wikidataId/Q34366'),
    ('6202010.xlsx', 'wikidataId/Q3235'),
    ('6202011.xlsx', 'wikidataId/Q3258'),
]

CSV_COLUMNS: List[str] = [
    'observationDate',
    'observationAbout',
    'Employed total ;  Persons ;',
    'Employed total ;  > Males ;',
    'Employed total ;  > Females ;',
    'Unemployed total ;  Persons ;',
    'Unemployed total ;  > Males ;',
    'Unemployed total ;  > Females ;',
    'Unemployment rate ;  Persons ;',
    'Unemployment rate ;  > Males ;',
    'Unemployment rate ;  > Females ;',
]

SERIES_COLUMNS: List[str] = CSV_COLUMNS[2:]
BASELINE_CUTOFF_DATE = '2023-01'


def create_retry_session(retries: int = 3,
                         backoff_factor: float = 2.0) -> requests.Session:
    """Creates a requests.Session configured with retry and backoff."""
    session = requests.Session()
    retry = Retry(
        total=retries,
        read=retries,
        connect=retries,
        backoff_factor=backoff_factor,
        status_forcelist=[429, 500, 502, 503, 504],
        allowed_methods=['HEAD', 'GET', 'OPTIONS'],
    )
    adapter = HTTPAdapter(max_retries=retry)
    session.mount('https://', adapter)
    session.mount('http://', adapter)
    return session


def resolve_latest_release_url(session: requests.Session,
                               explicit_url: str = '') -> str:
    """Resolves the latest ABS Labour Force monthly release URL."""
    if explicit_url:
        return explicit_url.rstrip('/')
    response = session.get(LATEST_RELEASE_URL, timeout=60)
    response.raise_for_status()
    soup = BeautifulSoup(response.text, 'html.parser')
    for link in soup.find_all('a', href=True):
        href = link['href']
        if '6202001.xlsx' in href:
            base_href = href.rsplit('/', 1)[0]
            if base_href.startswith('http'):
                return base_href
            return f'https://www.abs.gov.au{base_href}'

    candidates = []
    for match in re.finditer(
            r'/statistics/labour/employment-and-unemployment/'
            r'labour-force-australia/[a-z]{3}-\d{4}',
            response.text.lower(),
    ):
        cand = f'https://www.abs.gov.au{match.group(0)}'
        if cand not in candidates:
            candidates.append(cand)

    today = datetime.date.today()
    for offset in range(0, 24):
        year = today.year
        month = today.month - offset
        while month <= 0:
            month += 12
            year -= 1
        mon_abbr = calendar.month_abbr[month].lower()
        cand = f'{BASE_RELEASE_PREFIX}/{mon_abbr}-{year}'
        if cand not in candidates:
            candidates.append(cand)

    for candidate in candidates:
        probe = session.head(f'{candidate}/6202001.xlsx', timeout=30)
        if probe.status_code == 200:
            return candidate
    raise RuntimeError(
        f'Unable to resolve ABS Labour Force release URL from {LATEST_RELEASE_URL}'
    )


def format_series_value(raw_val: object, unit: str) -> str:
    """Formats an ABS Labour Force cell value according to its unit."""
    if raw_val is None or raw_val == '':
        return ''
    val = float(raw_val)
    if unit.strip() == '000':
        scaled = round(val * 1000.0, 4)
        text = f'{scaled:.4f}'.rstrip('0').rstrip('.')
        return text if text else '0'
    return f'{round(val, 1):.1f}'


def format_obs_date(raw_date: object) -> Optional[str]:
    """Formats a Data1 date cell into YYYY-MM."""
    if isinstance(raw_date, datetime.datetime):
        return raw_date.strftime('%Y-%m')
    if isinstance(raw_date, datetime.date):
        return raw_date.strftime('%Y-%m')
    if isinstance(raw_date, str):
        text = raw_date.strip()
        if re.match(r'^\d{4}-\d{2}$', text):
            return text
        if re.match(r'^\d{4}-\d{2}-\d{2}', text):
            return text[:7]
    return None


def parse_labour_force_workbook(workbook_bytes: bytes,
                                place_dcid: str) -> List[Dict[str, str]]:
    """Extracts Original monthly series from an ABS Labour Force workbook."""
    wb = openpyxl.load_workbook(io.BytesIO(workbook_bytes),
                                data_only=True,
                                read_only=True)
    if 'Data1' not in wb.sheetnames:
        raise ValueError(f'Missing Data1 sheet for {place_dcid}')
    ws = wb['Data1']
    rows = list(ws.iter_rows(values_only=True))
    if len(rows) < 11:
        raise ValueError(f'Insufficient rows in Data1 for {place_dcid}')

    header_titles = rows[0]
    header_units = rows[1]
    header_types = rows[2]

    series_col_map: Dict[str, Tuple[int, str]] = {}
    for col_idx in range(1, len(header_titles)):
        title = str(header_titles[col_idx] or '').strip()
        unit = str(header_units[col_idx] or '').strip()
        series_type = str(header_types[col_idx] or '').strip()
        if series_type == 'Original' and title in SERIES_COLUMNS:
            if title not in series_col_map:
                series_col_map[title] = (col_idx, unit)

    missing = [col for col in SERIES_COLUMNS if col not in series_col_map]
    if missing:
        raise ValueError(
            f'Missing Original series columns for {place_dcid}: {missing}')

    extracted: List[Dict[str, str]] = []
    for row in rows[10:]:
        if not row:
            continue
        obs_date = format_obs_date(row[0])
        if not obs_date:
            continue
        record: Dict[str, str] = {
            'observationDate': obs_date,
            'observationAbout': place_dcid,
        }
        has_all = True
        for col_name in SERIES_COLUMNS:
            col_idx, unit = series_col_map[col_name]
            val = row[col_idx] if col_idx < len(row) else None
            if val is None or val == '':
                has_all = False
                break
            record[col_name] = format_series_value(val, unit)
        if has_all:
            extracted.append(record)
    return extracted


def merge_employment_records(
    existing_rows: List[Dict[str, str]],
    new_rows: List[Dict[str, str]],
    cutoff_date: str = BASELINE_CUTOFF_DATE,
) -> List[Dict[str, str]]:
    """Preserves baseline rows <= cutoff_date and merges new rows > cutoff_date."""
    place_order = {place: idx for idx, (_, place) in enumerate(WORKBOOK_PLACES)}
    merged_by_key: Dict[Tuple[str, str], Dict[str, str]] = {}

    for row in existing_rows:
        place = row['observationAbout'].strip()
        obs_date = row['observationDate'].strip()
        merged_by_key[(place, obs_date)] = row

    for row in new_rows:
        place = row['observationAbout'].strip()
        obs_date = row['observationDate'].strip()
        if obs_date > cutoff_date or (place, obs_date) not in merged_by_key:
            merged_by_key[(place, obs_date)] = row

    sorted_keys = sorted(
        merged_by_key.keys(),
        key=lambda k: (place_order.get(k[0], 999), k[1]),
    )
    return [merged_by_key[k] for k in sorted_keys]


def refresh_labour_force_csv(
    csv_path: str,
    release_url: str = '',
    session: Optional[requests.Session] = None,
) -> int:
    """Downloads ABS Labour Force workbooks and updates aus_employment_int.csv."""
    own_session = session is None
    if own_session:
        session = create_retry_session()
    try:
        resolved_base = resolve_latest_release_url(session, release_url)
        logging.info('Resolved ABS Labour Force release URL: %s', resolved_base)

        existing_rows: List[Dict[str, str]] = []
        if os.path.exists(csv_path):
            with open(csv_path, 'r', encoding='utf-8') as f:
                reader = csv.DictReader(f)
                existing_rows = list(reader)

        new_rows: List[Dict[str, str]] = []
        for workbook_name, place_dcid in WORKBOOK_PLACES:
            url = f'{resolved_base}/{workbook_name}'
            logging.info('Downloading %s for %s', url, place_dcid)
            resp = session.get(url, timeout=60)
            resp.raise_for_status()
            place_rows = parse_labour_force_workbook(resp.content, place_dcid)
            logging.info(
                'Extracted %d monthly rows for %s (latest: %s)',
                len(place_rows),
                place_dcid,
                place_rows[-1]['observationDate'] if place_rows else 'N/A',
            )
            new_rows.extend(place_rows)

        merged_rows = merge_employment_records(existing_rows, new_rows)
        with open(csv_path, 'w', encoding='utf-8', newline='') as f:
            writer = csv.DictWriter(f, fieldnames=CSV_COLUMNS)
            writer.writeheader()
            writer.writerows(merged_rows)
        logging.info(
            'Wrote %d total rows to %s (was %d rows)',
            len(merged_rows),
            csv_path,
            len(existing_rows),
        )
        return len(merged_rows)
    finally:
        if own_session and session is not None:
            session.close()


def main(_):
    refresh_labour_force_csv(FLAGS.employment_csv, FLAGS.release_url)


if __name__ == '__main__':
    app.run(main)
