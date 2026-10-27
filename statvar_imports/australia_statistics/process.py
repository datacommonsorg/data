#!/usr/bin/env python3
# Copyright 2026 Google LLC
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#      http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.
"""Preprocesses raw ABS source files and runs stat_var_processor.py in parallel."""

from collections import defaultdict
import concurrent.futures
import csv
import glob
import io
import os
import re
import shutil
import subprocess
import sys
import tempfile
from typing import Any

from absl import app
from absl import flags
from absl import logging
import openpyxl
import xlrd

_MODULE_DIR = os.path.dirname(os.path.abspath(__file__))
_REPO_ROOT = os.path.abspath(os.path.join(_MODULE_DIR, '../..'))
_UTIL_DIR = os.path.join(_REPO_ROOT, 'util')
for _p in (_REPO_ROOT, _UTIL_DIR):
    if _p not in sys.path:
        sys.path.insert(0, _p)

import file_util

_PROCESSOR_SCRIPT = os.path.abspath(
    os.path.join(_MODULE_DIR,
                 '../../tools/statvar_importer/stat_var_processor.py'))

_FLAGS = flags.FLAGS
flags.DEFINE_string(
    'input_data',
    os.path.join(_MODULE_DIR, 'input_files/*'),
    'Comma-separated list or glob pattern of raw ABS input files.',
)
flags.DEFINE_string(
    'pv_map',
    os.path.join(_MODULE_DIR, 'australia_statistics_pv_map.csv'),
    'Path to the PV map CSV file.',
)
flags.DEFINE_string(
    'config_file',
    os.path.join(_MODULE_DIR, 'australia_statistics_metadata.csv'),
    'Path to the metadata config CSV file.',
)
flags.DEFINE_string(
    'output_path',
    os.path.join(_MODULE_DIR, 'output_files/australia_statistics_output'),
    'Output prefix path for merged CSV, TMCF, and MCF files.',
)
flags.DEFINE_string(
    'output_counters',
    os.path.join(_MODULE_DIR, 'counters/australia_statistics_counters.csv'),
    'Output path for merged counters CSV file.',
)
flags.DEFINE_string(
    'existing_statvar_mcf',
    'gs://unresolved_mcf/scripts/statvar/stat_vars.mcf',
    'Path to existing StatVar MCF file.',
)
flags.DEFINE_integer(
    'parallelism',
    2,
    'Maximum number of concurrent stat_var_processor worker processes.',
)

CSV_HEADER = [
    'observationDate',
    'indicatorCd',
    'regionCd',
    'value',
    'unit',
    'scalingFactor',
]

SDMX_REGION_CODES = {'AUS', '1', '2', '3', '4', '5', '6', '7', '8'}

STATE_NAMES = {
    'New South Wales',
    'Victoria',
    'Queensland',
    'South Australia',
    'Western Australia',
    'Tasmania',
    'Northern Territory',
    'Australian Capital Territory',
    'Capital Territory',
}

LF_MEASURES = {
    ('M3', '3'),
    ('M3', '1'),
    ('M3', '2'),
    ('M6', '3'),
    ('M6', '1'),
    ('M6', '2'),
    ('M13', '3'),
    ('M13', '1'),
    ('M13', '2'),
}

ERP_AGES = {
    'A04',
    'A59',
    'A10',
    'A15',
    'A20',
    'A25',
    'A30',
    'A35',
    'A40',
    'A45',
    'A50',
    'A55',
    'A60',
    'A65',
    'A70',
    'A75',
    'A80',
    '8599',
    'TOT',
}

HIST_AGE_LABELS = {
    '0-4',
    '0-4 years',
    '5-9',
    '5-9 years',
    '10-14',
    '10-14 years',
    '15-19',
    '15-19 years',
    '20-24',
    '20-24 years',
    '25-29',
    '25-29 years',
    '30-34',
    '30-34 years',
    '35-39',
    '35-39 years',
    '40-44',
    '40-44 years',
    '45-49',
    '45-49 years',
    '50-54',
    '50-54 years',
    '55-59',
    '55-59 years',
    '60-64',
    '60-64 years',
    '65-69',
    '65-69 years',
    '70-74',
    '70-74 years',
    '75-79',
    '75-79 years',
    '80-84',
    '80-84 years',
    '85 and over',
    '85 years and over',
    'Total',
}

HEALTH_CONDITIONS = {
    'Arthritis',
    'Asthma',
    'Cancer (including remission)',
    'Dementia (including Alzheimer\u2019s)',
    "Dementia (including Alzheimer's)",
    'Dementia (including Alzheimer)',
    'Diabetes (excluding gestational diabetes)',
    'Heart disease (including heart attack or angina)',
    'Kidney disease',
    'Lung condition (including COPD or emphysema)(d)',
    'Lung condition (including COPD or emphysema)(e)',
    'Mental health condition (including depression or anxiety)',
    'Stroke',
}

INCOME_LABELS = {
    'Lowest quintile',
    'Second quintile',
    'Third quintile',
    'Fourth quintile',
    'Highest quintile',
    '10th (P10)',
    '20th (P20)',
    '30th (P30)',
    '40th (P40)',
    '50th (P50)',
    '60th (P60)',
    '70th (P70)',
    '80th (P80)',
    '90th (P90)',
}


def _format_number(val: Any, precision: int = 4) -> str:
    """Formats numeric values cleanly without trailing zeroes."""
    f = float(val)
    if f.is_integer():
        return str(int(round(f)))
    return str(round(f, precision)).rstrip('0').rstrip('.')


def _make_row(
    observation_date: str,
    indicator_cd: str,
    region_cd: str,
    value: str,
    unit: str = '',
    scaling_factor: str = '',
) -> dict[str, str]:
    """Constructs a standardized input CSV row dictionary."""
    return {
        'observationDate': str(observation_date),
        'indicatorCd': str(indicator_cd),
        'regionCd': str(region_cd),
        'value': str(value),
        'unit': str(unit),
        'scalingFactor': str(scaling_factor),
    }


def parse_sdmx_labour_force_csv(csv_text: str) -> list[dict[str, str]]:
    """Extracts raw ABS SDMX Labour Force rows without scaling."""
    rows: list[dict[str, str]] = []
    reader = csv.DictReader(io.StringIO(csv_text))
    for r in reader:
        region = r.get('REGION', '').strip()
        measure = r.get('MEASURE', '').strip()
        sex = r.get('SEX', '').strip()
        period = r.get('TIME_PERIOD', '').strip()
        obs_val = r.get('OBS_VALUE', '').strip()
        if (not obs_val or region not in SDMX_REGION_CODES
                or (measure, sex) not in LF_MEASURES):
            continue
        indicator_cd = f'LF_{measure}_{sex}'
        if measure == 'M13':
            val_str = _format_number(float(obs_val), precision=1)
        else:
            val_str = obs_val
        rows.append(_make_row(period, indicator_cd, region, val_str))
    return rows


def parse_sdmx_erp_q_csv(csv_text: str) -> list[dict[str, str]]:
    """Extracts raw ABS SDMX Quarterly ERP rows (June Q2 annual reference)."""
    rows: list[dict[str, str]] = []
    reader = csv.DictReader(io.StringIO(csv_text))
    for r in reader:
        period = r.get('TIME_PERIOD', '').strip()
        if not period.endswith('-Q2'):
            continue
        year = period.split('-')[0]
        if year == '2021':
            # 2021 population is sourced from the 2021 Census Population Data Summary.
            continue
        region = r.get('REGION', '').strip()
        sex = r.get('SEX', '').strip()
        age = r.get('AGE', '').strip()
        obs_val = r.get('OBS_VALUE', '').strip()
        if (not obs_val or region not in SDMX_REGION_CODES
                or sex not in ('1', '2', '3') or age not in ERP_AGES):
            continue
        indicator_cd = f'ERP_{sex}_{age}'
        val_str = str(int(round(float(obs_val))))
        rows.append(_make_row(year, indicator_cd, region, val_str))
    return rows


def parse_census_2021_population_xlsx(
    xlsx_bytes: bytes, ) -> tuple[list[dict[str, str]], list[dict[str, str]]]:
    """Parses 2021 Census Population data summary.xlsx (Table 1 and Table 3)."""
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes), data_only=True)
    gender_rows: list[dict[str, str]] = []
    ws1 = wb['Table 1']
    for r in range(1, ws1.max_row + 1):
        state_name = str(ws1.cell(r, 1).value or '').strip()
        if state_name in STATE_NAMES:
            male_v = ws1.cell(r, 2).value
            female_v = ws1.cell(r, 3).value
            total_v = ws1.cell(r, 4).value
            if male_v is not None:
                gender_rows.append(
                    _make_row('2021', 'POP_Males_Total', state_name,
                              str(int(male_v))))
            if female_v is not None:
                gender_rows.append(
                    _make_row(
                        '2021',
                        'POP_Females_Total',
                        state_name,
                        str(int(female_v)),
                    ))
            if total_v is not None:
                gender_rows.append(
                    _make_row(
                        '2021',
                        'POP_Persons_Total',
                        state_name,
                        str(int(total_v)),
                    ))

    agegrp_rows: list[dict[str, str]] = []
    ws3 = wb['Table 3']
    for r in range(1, ws3.max_row + 1):
        age_lbl = str(ws3.cell(r, 1).value or '').strip()
        if age_lbl in HIST_AGE_LABELS:
            male_v = ws3.cell(r, 2).value
            female_v = ws3.cell(r, 3).value
            total_v = ws3.cell(r, 4).value
            if male_v is not None:
                agegrp_rows.append(
                    _make_row(
                        '2021',
                        f'POP_Males_{age_lbl}',
                        'AUS_DBL',
                        str(int(male_v)),
                    ))
            if female_v is not None:
                agegrp_rows.append(
                    _make_row(
                        '2021',
                        f'POP_Females_{age_lbl}',
                        'AUS_DBL',
                        str(int(female_v)),
                    ))
            if total_v is not None:
                agegrp_rows.append(
                    _make_row(
                        '2021',
                        f'POP_Persons_{age_lbl}',
                        'AUS_DBL',
                        str(int(total_v)),
                    ))
    return gender_rows, agegrp_rows


def parse_historical_population_xls(
    xls_bytes: bytes,
) -> tuple[list[dict[str, str]], list[dict[str, str]], list[dict[str, str]]]:
    """Parses ABS 3105065001ds0002_2019.xls historical population tables."""
    wb = xlrd.open_workbook(file_contents=xls_bytes)

    def _parse_sheet(sheet_name: str,
                     region_cd: str,
                     five_to_six: bool = False) -> list[dict[str, str]]:
        sh = wb.sheet_by_name(sheet_name)
        years = {
            c: str(int(float(sh.cell_value(4, c))))
            for c in range(1, sh.ncols) if sh.cell_value(4, c) != ''
        }
        gender_lbl: str | None = None
        out: list[dict[str, str]] = []
        for r in range(5, sh.nrows):
            lbl = str(sh.cell_value(r, 0)).strip()
            if lbl in ('Males', 'Females', 'Persons'):
                gender_lbl = lbl
                continue
            if lbl in HIST_AGE_LABELS and gender_lbl is not None:
                age_key = ('6-9' if (five_to_six
                                     and lbl in ('5-9', '5-9 years')) else lbl)
                indicator_cd = f'POP_{gender_lbl}_{age_key}'
                for c, yr in years.items():
                    val = sh.cell_value(r, c)
                    if val != '' and val is not None:
                        out.append(
                            _make_row(
                                yr,
                                indicator_cd,
                                region_cd,
                                str(int(round(float(val)))),
                            ))
        return out

    aus_all = _parse_sheet('Table 2.1', 'AUS', five_to_six=False)
    aus_1901_1985 = [r for r in aus_all if int(r['observationDate']) <= 1985]
    aus_1986_2016 = [r for r in aus_all if int(r['observationDate']) >= 1986]

    state_sheets = [
        ('Table 2.4', 'NSW_DBL'),
        ('Table 2.6', 'VIC_DBL'),
        ('Table 2.8', 'QLD_DBL'),
        ('Table 2.10', 'SA_DBL'),
        ('Table 2.12', 'WA_DBL'),
        ('Table 2.15', 'TAS_DBL'),
    ]
    states_rows: list[dict[str, str]] = []
    for sname, reg_cd in state_sheets:
        states_rows.extend(_parse_sheet(sname, reg_cd, five_to_six=True))

    return aus_1901_1985, aus_1986_2016, states_rows


def parse_health_2021_xlsx(
    xlsx_bytes: bytes, ) -> tuple[list[dict[str, str]], list[dict[str, str]]]:
    """Parses 2021 Census Health data summary.xlsx (Table 1 and Table 3)."""
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes), data_only=True)

    ws1 = wb['Table 1']
    col_ages: dict[int, str] = {}
    for c in range(2, 21):
        top = str(ws1.cell(11, c).value or '').strip()
        bot = str(ws1.cell(12, c).value or '').strip()
        combined = f'{top} {bot}'.strip()
        if combined.startswith('Total'):
            col_ages[c] = 'Total'
        elif combined in HIST_AGE_LABELS:
            col_ages[c] = combined

    national_rows: list[dict[str, str]] = []
    sec_gender: str | None = None
    for r in range(13, ws1.max_row + 1):
        sec_hdr = str(ws1.cell(r, 2).value or '').strip()
        if sec_hdr in ('MALES', 'FEMALES', 'PERSONS'):
            sec_gender = sec_hdr
            continue
        cond_lbl = str(ws1.cell(r, 1).value or '').strip()
        if cond_lbl in HEALTH_CONDITIONS and sec_gender is not None:
            for c, age_lbl in col_ages.items():
                val = ws1.cell(r, c).value
                if val is None:
                    continue
                if age_lbl == 'Total' and sec_gender == 'PERSONS':
                    continue
                indicator_cd = f'HEALTH_{sec_gender}_{age_lbl}_{cond_lbl}'
                national_rows.append(
                    _make_row('2021', indicator_cd, 'AUS', str(int(val))))

    ws3 = wb['Table 3']
    col_regions = {
        2: 'NSW',
        3: 'VIC',
        4: 'QLD',
        5: 'SA',
        6: 'WA',
        7: 'TAS',
        8: 'NT',
        9: 'ACT',
        10: 'AUS',
    }
    state_rows: list[dict[str, str]] = []
    for r in range(13, ws3.max_row + 1):
        cond_lbl = str(ws3.cell(r, 1).value or '').strip()
        if cond_lbl in HEALTH_CONDITIONS:
            indicator_cd = f'HEALTH_Total_{cond_lbl}'
            for c, reg_cd in col_regions.items():
                val = ws3.cell(r, c).value
                if val is not None:
                    state_rows.append(
                        _make_row('2021', indicator_cd, reg_cd, str(int(val))))

    return national_rows, state_rows


def parse_household_income_xlsx(xlsx_bytes: bytes) -> list[dict[str, str]]:
    """Parses ABS Household income and income distribution Table 1.1."""
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes), data_only=True)
    ws = wb['Table 1.1']
    col_years: dict[int, str] = {}
    for c in range(3, ws.max_column + 1):
        hdr = str(ws.cell(5, c).value or '').strip()
        if hdr:
            norm = hdr.replace('\u2013', '-').replace('\u2014', '-')
            if '-' in norm:
                start_part, end_part = norm.split('-', 1)
                year = (start_part[:2] +
                        end_part if len(end_part) == 2 else end_part)
                col_years[c] = year

    rows: list[dict[str, str]] = []
    in_income_section = True
    for r in range(6, ws.max_row + 1):
        lbl = str(ws.cell(r, 1).value or '').strip()
        if lbl == 'Income share':
            in_income_section = False
            break
        if in_income_section and lbl in INCOME_LABELS:
            indicator_cd = f'INC_{lbl}'
            for c, yr in col_years.items():
                val = ws.cell(r, c).value
                if val is not None:
                    rows.append(
                        _make_row(
                            yr,
                            indicator_cd,
                            'AUS',
                            _format_number(val),
                            unit='AUD',
                        ))
    return rows


def parse_historical_life_expectancy_xls(
    xls_bytes: bytes, ) -> tuple[list[dict[str, str]], list[dict[str, str]]]:
    """Parses ABS 3105065001ds0006_2019.xls (Tables 6.1, 6.2, and 6.6)."""
    wb = xlrd.open_workbook(file_contents=xls_bytes)

    sh1 = wb.sheet_by_name('Table 6.1')
    valid_places = {
        'NSW',
        'Vic.',
        'Qld',
        'SA',
        'WA',
        'Tas.',
        'NT',
        'ACT(a)',
        'Australia(b)',
    }
    state_rows: list[dict[str, str]] = []
    gender_lbl: str | None = None
    for r in range(5, sh1.nrows):
        lbl = str(sh1.cell_value(r, 0)).strip()
        if lbl in ('Males', 'Females'):
            gender_lbl = lbl
            continue
        if lbl in valid_places and gender_lbl is not None:
            indicator_cd = f'LE_{gender_lbl}_0'
            for c in range(10, sh1.ncols):
                col_hdr = str(sh1.cell_value(4, c)).strip()
                if not col_hdr or col_hdr == '1993-1995':
                    continue
                yr = (col_hdr.split('-')[1] if '-' in col_hdr else str(
                    int(float(col_hdr))))
                val = sh1.cell_value(r, c)
                if val != '' and val is not None:
                    state_rows.append(
                        _make_row(
                            yr,
                            indicator_cd,
                            lbl,
                            _format_number(val, precision=2),
                        ))

    age_rows: list[dict[str, str]] = []
    for sname, g_lbl in [('Table 6.2', 'Males'), ('Table 6.6', 'Females')]:
        sh = wb.sheet_by_name(sname)
        cols = {
            c: str(sh.cell_value(4, c)).strip().split('-')[1]
            for c in range(1, sh.ncols) if str(sh.cell_value(4, c)).strip()
        }
        for r in range(6, sh.nrows):
            age_v = sh.cell_value(r, 0)
            if age_v == '' or not isinstance(age_v, (int, float)):
                continue
            age = int(round(age_v))
            if 1 <= age <= 99:
                indicator_cd = f'LE_{g_lbl}_{age}'
                for c, yr in cols.items():
                    val = sh.cell_value(r, c)
                    if val != '' and val is not None:
                        age_rows.append(
                            _make_row(
                                yr,
                                indicator_cd,
                                'AUS',
                                _format_number(val, precision=2),
                            ))

    return age_rows, state_rows


def parse_recent_life_expectancy_workbook(
        file_bytes: bytes, ext: str,
        end_year: str) -> tuple[list[dict[str, str]], list[dict[str, str]]]:
    """Parses a recent ABS Life Tables workbook (3302055001DO001_YYYYYYYY)."""
    table_idx_to_region = {
        1: 'NSW',
        2: 'VIC',
        3: 'QLD',
        4: 'SA',
        5: 'WA',
        6: 'TAS',
        7: 'NT',
        8: 'ACT',
        9: 'AUS',
    }
    age_rows: list[dict[str, str]] = []
    state_rows: list[dict[str, str]] = []

    if ext == 'xls':
        wb = xlrd.open_workbook(file_contents=file_bytes)
        for idx, reg_cd in table_idx_to_region.items():
            sname = (f'Table_1.{idx}' if f'Table_1.{idx}' in wb.sheet_names()
                     else f'Table {idx}')
            sh = wb.sheet_by_name(sname)
            for r in range(6, sh.nrows):
                age_v = sh.cell_value(r, 0)
                if age_v == '' or not isinstance(age_v, (int, float)):
                    continue
                age = int(round(age_v))
                male_ex = sh.cell_value(r, 4)
                female_ex = sh.cell_value(r, 8)
                if age == 0:
                    if male_ex != '':
                        state_rows.append(
                            _make_row(
                                end_year,
                                'LE_Males_0',
                                reg_cd,
                                _format_number(male_ex, precision=2),
                            ))
                    if female_ex != '':
                        state_rows.append(
                            _make_row(
                                end_year,
                                'LE_Females_0',
                                reg_cd,
                                _format_number(female_ex, precision=2),
                            ))
                elif reg_cd == 'AUS' and 1 <= age <= 99:
                    if male_ex != '':
                        age_rows.append(
                            _make_row(
                                end_year,
                                f'LE_Males_{age}',
                                'AUS',
                                _format_number(male_ex, precision=2),
                            ))
                    if female_ex != '':
                        age_rows.append(
                            _make_row(
                                end_year,
                                f'LE_Females_{age}',
                                'AUS',
                                _format_number(female_ex, precision=2),
                            ))
    else:
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
        for idx, reg_cd in table_idx_to_region.items():
            sname = (f'Table_1.{idx}'
                     if f'Table_1.{idx}' in wb.sheetnames else f'Table {idx}')
            sh = wb[sname]
            for r in range(6, sh.max_row + 1):
                age_v = sh.cell(r, 1).value
                if age_v is None or not isinstance(age_v, (int, float)):
                    continue
                age = int(round(age_v))
                male_ex = sh.cell(r, 5).value
                female_ex = sh.cell(r, 9).value
                if age == 0:
                    if male_ex is not None:
                        state_rows.append(
                            _make_row(
                                end_year,
                                'LE_Males_0',
                                reg_cd,
                                _format_number(male_ex, precision=2),
                            ))
                    if female_ex is not None:
                        state_rows.append(
                            _make_row(
                                end_year,
                                'LE_Females_0',
                                reg_cd,
                                _format_number(female_ex, precision=2),
                            ))
                elif reg_cd == 'AUS' and 1 <= age <= 99:
                    if male_ex is not None:
                        age_rows.append(
                            _make_row(
                                end_year,
                                f'LE_Males_{age}',
                                'AUS',
                                _format_number(male_ex, precision=2),
                            ))
                    if female_ex is not None:
                        age_rows.append(
                            _make_row(
                                end_year,
                                f'LE_Females_{age}',
                                'AUS',
                                _format_number(female_ex, precision=2),
                            ))
    return age_rows, state_rows


def _atomic_write_text(target_path: str, content: str) -> None:
    """Writes text content atomically to target_path via a temp file."""
    target_dir = os.path.dirname(os.path.abspath(target_path))
    os.makedirs(target_dir, exist_ok=True)
    fd, tmp_path = tempfile.mkstemp(
        prefix='.tmp_proc_',
        suffix=os.path.splitext(target_path)[1],
        dir=target_dir,
    )
    try:
        with os.fdopen(fd, 'w', encoding='utf-8', newline='') as f:
            f.write(content)
        if not (os.path.exists(tmp_path) and os.path.getsize(tmp_path) > 0):
            raise RuntimeError(
                f'Atomic write failed: {tmp_path} is empty or missing.')
        os.replace(tmp_path, target_path)
    except Exception:
        if os.path.exists(tmp_path):
            os.remove(tmp_path)
        raise


def _write_staged_csv(target_path: str, rows: list[dict[str, str]]) -> str:
    """Writes staged CSV rows to target_path and returns target_path."""
    if not rows:
        raise RuntimeError(f'No rows produced for staged file: {target_path}')
    buf = io.StringIO()
    writer = csv.DictWriter(buf, fieldnames=CSV_HEADER, lineterminator='\n')
    writer.writeheader()
    writer.writerows(rows)
    _atomic_write_text(target_path, buf.getvalue())
    return target_path


def _expand_input_files(input_data: str) -> list[str]:
    """Expands comma-separated paths or glob patterns into non-empty file paths."""
    files: list[str] = []
    for pattern in input_data.split(','):
        pattern = pattern.strip()
        if not pattern:
            continue
        matched = sorted(glob.glob(pattern))
        if matched:
            files.extend(matched)
        elif os.path.exists(pattern):
            files.append(pattern)
    valid_files = [
        p for p in set(files) if os.path.isfile(p) and os.path.getsize(p) > 0
    ]
    return sorted(valid_files, key=lambda p: -os.path.getsize(p))


def _stage_inputs(raw_files: list[str], staged_dir: str) -> list[str]:
    """Converts raw ABS SDMX CSVs and Excel workbooks into staged CSV shards."""
    os.makedirs(staged_dir, exist_ok=True)
    by_name = {os.path.basename(p): p for p in raw_files}
    staged_files: list[str] = []

    # 1. Pass through already-standardized CSVs (e.g. test_data/sample_input.csv)
    for fname, fpath in sorted(by_name.items()):
        if fname.endswith('.csv') and fname not in (
                'abs_lf_sdmx.csv',
                'abs_erp_q_sdmx.csv',
        ):
            with open(fpath, 'r', encoding='utf-8-sig') as f:
                first_line = f.readline().strip()
            if first_line.startswith('observationDate,'):
                staged_files.append(fpath)

    # 2. Labour Force SDMX CSV
    if 'abs_lf_sdmx.csv' in by_name:
        with open(by_name['abs_lf_sdmx.csv'], 'r', encoding='utf-8') as f:
            lf_rows = parse_sdmx_labour_force_csv(f.read())
        staged_files.append(
            _write_staged_csv(os.path.join(staged_dir, 'aus_employment.csv'),
                              lf_rows))

    # 3. 2021 Census Population Excel workbook
    if 'population_data_summary_2021.xlsx' in by_name:
        with open(by_name['population_data_summary_2021.xlsx'], 'rb') as f:
            g_2021, a_2021 = parse_census_2021_population_xlsx(f.read())
        staged_files.append(
            _write_staged_csv(
                os.path.join(staged_dir, 'aus_count_person_gender_2021.csv'),
                g_2021,
            ))
        staged_files.append(
            _write_staged_csv(
                os.path.join(staged_dir,
                             'aus_count_person_gender_agegrp_2021.csv'),
                a_2021,
            ))

    # 4. Historical Population (.xls) + Quarterly ERP SDMX (.csv)
    aus_1901_1985: list[dict[str, str]] = []
    aus_1986_plus: list[dict[str, str]] = []
    states_combined: list[dict[str, str]] = []
    if 'historical_population_2016.xls' in by_name:
        with open(by_name['historical_population_2016.xls'], 'rb') as f:
            p1, p2, p_states = parse_historical_population_xls(f.read())
        aus_1901_1985.extend(p1)
        aus_1986_plus.extend(p2)
        states_combined.extend(p_states)
    if 'abs_erp_q_sdmx.csv' in by_name:
        with open(by_name['abs_erp_q_sdmx.csv'], 'r', encoding='utf-8') as f:
            erp_rows = parse_sdmx_erp_q_csv(f.read())
        aus_1986_plus.extend([r for r in erp_rows if r['regionCd'] == 'AUS'])
        states_combined.extend([r for r in erp_rows if r['regionCd'] != 'AUS'])

    if aus_1901_1985:
        staged_files.append(
            _write_staged_csv(
                os.path.join(
                    staged_dir,
                    'aus_count_person_gender_agegrp_historic_1901to1985.csv',
                ),
                aus_1901_1985,
            ))
    if aus_1986_plus:
        staged_files.append(
            _write_staged_csv(
                os.path.join(
                    staged_dir,
                    'aus_count_person_gender_agegrp_historic_1986to2016.csv',
                ),
                aus_1986_plus,
            ))
    if states_combined:
        staged_files.append(
            _write_staged_csv(
                os.path.join(
                    staged_dir,
                    'aus_count_person_gender_agegrp_historic_states.csv',
                ),
                states_combined,
            ))

    # 5. 2021 Census Health Conditions Excel workbook
    if 'health_data_summary_2021.xlsx' in by_name:
        with open(by_name['health_data_summary_2021.xlsx'], 'rb') as f:
            nat_health, state_health = parse_health_2021_xlsx(f.read())
        staged_files.append(
            _write_staged_csv(
                os.path.join(staged_dir, 'aus_health_condition.csv'),
                nat_health,
            ))
        staged_files.append(
            _write_staged_csv(
                os.path.join(staged_dir, 'aus_health_condition_state.csv'),
                state_health,
            ))

    # 6. Historical + Recent Life Expectancy workbooks
    le_age_rows: list[dict[str, str]] = []
    le_state_rows: list[dict[str, str]] = []
    if 'historical_life_expectancy_2016.xls' in by_name:
        with open(by_name['historical_life_expectancy_2016.xls'], 'rb') as f:
            h_age, h_state = parse_historical_life_expectancy_xls(f.read())
        le_age_rows.extend(h_age)
        le_state_rows.extend(h_state)

    le_pattern = re.compile(r'^life_expectancy_\d{4}(\d{4})\.(xlsx|xls)$')
    for fname in sorted(by_name.keys()):
        m = le_pattern.match(fname)
        if m:
            end_yr, ext = m.group(1), m.group(2)
            with open(by_name[fname], 'rb') as f:
                r_age, r_state = parse_recent_life_expectancy_workbook(
                    f.read(), ext, end_yr)
            le_age_rows.extend(r_age)
            le_state_rows.extend(r_state)

    if le_age_rows:
        staged_files.append(
            _write_staged_csv(
                os.path.join(staged_dir, 'aus_life_expectancy_gender_age.csv'),
                le_age_rows,
            ))
    if le_state_rows:
        staged_files.append(
            _write_staged_csv(
                os.path.join(staged_dir, 'aus_life_expectancy_states.csv'),
                le_state_rows,
            ))

    # 7. Weekly Household Income Excel workbook
    if 'household_income_2019_20.xlsx' in by_name:
        with open(by_name['household_income_2019_20.xlsx'], 'rb') as f:
            inc_rows = parse_household_income_xlsx(f.read())
        staged_files.append(
            _write_staged_csv(
                os.path.join(staged_dir, 'aus_weekly_household_income.csv'),
                inc_rows,
            ))

    return sorted(staged_files, key=lambda p: -os.path.getsize(p))


def _prepare_local_statvar_mcf(existing_statvar_mcf: str,
                               shards_tmp_dir: str) -> str:
    """Copies GCS existing_statvar_mcf locally once so shards do not re-download."""
    if not existing_statvar_mcf or not existing_statvar_mcf.startswith(
            'gs://'):
        return existing_statvar_mcf
    local_mcf = os.path.join(shards_tmp_dir, 'cached_stat_vars.mcf')
    logging.info('Caching %s locally to %s...', existing_statvar_mcf,
                 local_mcf)
    file_util.file_copy(existing_statvar_mcf, local_mcf)
    return local_mcf


def _run_shard(
    shard_idx: int,
    input_file: str,
    shard_output_prefix: str,
    shard_counters_path: str,
    pv_map: str,
    config_file: str,
    existing_statvar_mcf: str,
) -> tuple[int, str, str, str]:
    """Runs stat_var_processor.py on a single staged CSV in an isolated work dir."""
    shard_work_dir = f'{shard_output_prefix}_cwd'
    os.makedirs(shard_work_dir, exist_ok=True)
    cmd = [
        sys.executable,
        _PROCESSOR_SCRIPT,
        f'--input_data={input_file}',
        f'--pv_map={pv_map}',
        f'--config_file={config_file}',
        f'--output_path={shard_output_prefix}',
        f'--output_counters={shard_counters_path}',
        f'--existing_statvar_mcf={existing_statvar_mcf}',
    ]
    logging.info('[Shard %02d] Starting %s', shard_idx,
                 os.path.basename(input_file))
    proc = subprocess.run(cmd,
                          cwd=shard_work_dir,
                          capture_output=True,
                          text=True)
    if proc.returncode != 0:
        raise RuntimeError(
            f'Shard {shard_idx} ({input_file}) failed with exit code {proc.returncode}:\n'
            f'{proc.stderr[-2000:]}')
    logging.info('[Shard %02d] Finished %s', shard_idx,
                 os.path.basename(input_file))
    return shard_idx, input_file, shard_output_prefix, shard_counters_path


def _merge_csv_files(shard_prefixes: list[str], final_csv_path: str) -> int:
    """Merges shard CSV files into final_csv_path, deduplicating rows."""
    seen_lines: set[str] = set()
    ordered_lines: list[str] = []
    header_line = ''

    for prefix in shard_prefixes:
        shard_csv = f'{prefix}.csv'
        if not os.path.exists(shard_csv):
            continue
        with open(shard_csv, 'r', encoding='utf-8') as in_f:
            header = in_f.readline()
            if not header:
                continue
            if not header_line:
                header_line = header
            for line in in_f:
                if line and line not in seen_lines:
                    seen_lines.add(line)
                    ordered_lines.append(line)

    if not ordered_lines:
        raise RuntimeError(
            f'Merged CSV produced 0 data rows for {final_csv_path}')

    _atomic_write_text(final_csv_path, header_line + ''.join(ordered_lines))
    return len(ordered_lines)


def _counter_aggregation_mode(key: str) -> str:
    """Determines how a counter key should be aggregated across shards."""
    lower = key.lower()
    if (lower == 'existing-statvars-from-mcf' or lower.startswith('existing-')
            or lower.endswith('output-svobs-unique-typeof')):
        return 'max'
    if 'start_time' in lower:
        return 'min'
    if 'end_time' in lower or 'mem' in lower:
        return 'max'
    if lower.endswith('_rate') or '_rate_' in lower:
        return 'avg'
    return 'sum'


def _merge_counters(shard_counter_files: list[str],
                    final_counters_path: str) -> None:
    """Aggregates numeric counters across all shards into final_counters_path."""
    if not final_counters_path:
        return
    values_by_key: dict[str, list[float]] = defaultdict(list)
    is_int_map: dict[str, bool] = defaultdict(lambda: True)

    for cf in shard_counter_files:
        if not os.path.exists(cf):
            continue
        with open(cf, 'r', encoding='utf-8') as f:
            for line in f:
                line = line.strip()
                if not line or line.startswith('key,'):
                    continue
                parts = line.rsplit(',', 1)
                if len(parts) != 2:
                    continue
                k, v_str = parts[0].strip('"'), parts[1].strip()
                try:
                    if '.' in v_str:
                        values_by_key[k].append(float(v_str))
                        is_int_map[k] = False
                    else:
                        values_by_key[k].append(float(int(v_str)))
                except ValueError:
                    continue

    out_lines = ['key,value\n']
    for k in sorted(values_by_key.keys()):
        vals = values_by_key[k]
        mode = _counter_aggregation_mode(k)
        if mode == 'max':
            agg = max(vals)
        elif mode == 'min':
            agg = min(vals)
        elif mode == 'avg':
            agg = sum(vals) / len(vals)
            is_int_map[k] = False
        else:
            agg = sum(vals)

        if is_int_map[k]:
            out_lines.append(f'{k},{int(round(agg))}\n')
        else:
            out_lines.append(f'{k},{agg:.2f}\n')

    _atomic_write_text(final_counters_path, ''.join(out_lines))


def _merge_mcf_files(shard_prefixes: list[str], suffix: str,
                     final_mcf_path: str) -> None:
    """Deduplicates MCF nodes by Node line across shards and writes final_mcf_path."""
    seen_nodes: dict[str, str] = {}
    for prefix in shard_prefixes:
        shard_mcf = f'{prefix}{suffix}'
        if not os.path.exists(shard_mcf):
            continue
        with open(shard_mcf, 'r', encoding='utf-8') as f:
            content = f.read()
        for block in content.split('\n\n'):
            block = block.strip()
            if not block:
                continue
            lines = [
                l for l in block.splitlines() if not l.strip().startswith('#')
            ]
            if not lines:
                continue
            node_line = next((l for l in lines if l.startswith('Node:')),
                             lines[0])
            if node_line not in seen_nodes:
                seen_nodes[node_line] = '\n'.join(lines)

    if os.path.isdir(final_mcf_path):
        shutil.rmtree(final_mcf_path)
    elif os.path.exists(final_mcf_path):
        os.remove(final_mcf_path)

    if not seen_nodes:
        return

    body = ('\n\n'.join(seen_nodes[k]
                        for k in sorted(seen_nodes.keys())) + '\n\n')
    _atomic_write_text(final_mcf_path, body)


def _run_pipeline() -> None:
    raw_files = _expand_input_files(_FLAGS.input_data)
    if not raw_files:
        raise RuntimeError(f'No input files matched: {_FLAGS.input_data}')

    output_path = os.path.abspath(_FLAGS.output_path)
    output_dir = os.path.dirname(output_path)
    os.makedirs(output_dir, exist_ok=True)

    shards_tmp_dir = tempfile.mkdtemp(prefix='.shards_', dir=output_dir)
    try:
        staged_dir = os.path.join(shards_tmp_dir, 'staged_inputs')
        input_files = _stage_inputs(raw_files, staged_dir)
        if not input_files:
            raise RuntimeError(
                f'Staging produced 0 CSV files from: {raw_files}')

        local_statvar_mcf = _prepare_local_statvar_mcf(
            _FLAGS.existing_statvar_mcf, shards_tmp_dir)

        max_workers = max(
            1, min(_FLAGS.parallelism, len(input_files),
                   os.cpu_count() or 1))
        logging.info(
            'Processing %d staged CSV files across %d parallel workers...',
            len(input_files),
            max_workers,
        )

        shard_prefixes: list[str] = []
        shard_counters: list[str] = []
        futures = []

        with concurrent.futures.ThreadPoolExecutor(
                max_workers=max_workers) as executor:
            for idx, input_file in enumerate(input_files):
                base = os.path.splitext(os.path.basename(input_file))[0]
                s_prefix = os.path.join(shards_tmp_dir,
                                        f'shard_{idx:02d}_{base}')
                s_counters = os.path.join(shards_tmp_dir,
                                          f'counters_{idx:02d}_{base}.csv')
                shard_prefixes.append(s_prefix)
                shard_counters.append(s_counters)
                futures.append(
                    executor.submit(
                        _run_shard,
                        idx,
                        os.path.abspath(input_file),
                        s_prefix,
                        s_counters,
                        os.path.abspath(_FLAGS.pv_map),
                        os.path.abspath(_FLAGS.config_file),
                        local_statvar_mcf,
                    ))

            for fut in concurrent.futures.as_completed(futures):
                fut.result()

        # 1. Merge CSV outputs
        final_csv = f'{output_path}.csv'
        total_rows = _merge_csv_files(shard_prefixes, final_csv)
        logging.info('Merged %d SVObs rows into %s', total_rows, final_csv)

        # 2. Generate unified TMCF (normalizing table reference to output_path basename)
        final_tmcf = f'{output_path}.tmcf'
        table_name = os.path.basename(output_path)
        for prefix in shard_prefixes:
            shard_tmcf = f'{prefix}.tmcf'
            if os.path.exists(shard_tmcf):
                shard_table = os.path.basename(prefix)
                with open(shard_tmcf, 'r', encoding='utf-8') as f:
                    tmcf_text = (f.read().replace(f'E:{shard_table}->',
                                                  f'E:{table_name}->').replace(
                                                      f'C:{shard_table}->',
                                                      f'C:{table_name}->'))
                _atomic_write_text(final_tmcf, tmcf_text)
                break
        if not os.path.exists(final_tmcf):
            raise RuntimeError(f'Failed to generate TMCF file: {final_tmcf}')

        # 3. Merge StatVar MCF and Schema MCF files
        _merge_mcf_files(shard_prefixes, '_stat_vars.mcf',
                         f'{output_path}_stat_vars.mcf')
        _merge_mcf_files(
            shard_prefixes,
            '_stat_vars_schema.mcf',
            f'{output_path}_stat_vars_schema.mcf',
        )

        # 4. Merge Counters
        if _FLAGS.output_counters:
            _merge_counters(shard_counters,
                            os.path.abspath(_FLAGS.output_counters))
            logging.info('Merged counters into %s', _FLAGS.output_counters)

    finally:
        if os.path.exists(shards_tmp_dir):
            shutil.rmtree(shards_tmp_dir)


def main(_) -> None:
    try:
        _run_pipeline()
    except Exception as e:
        logging.fatal('%s', e, exc_info=True)


if __name__ == '__main__':
    app.run(main)
