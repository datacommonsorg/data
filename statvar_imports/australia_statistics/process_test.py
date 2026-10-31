#!/usr/bin/env python3
# Copyright 2026 Google LLC
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#      http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.
"""Unit tests for Australia Statistics process.py."""

import io
import os
import sys
import tempfile
import unittest
from unittest import mock

import openpyxl

# Ensure repository root is in sys.path for both local and root-level test runs
_SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
_REPO_ROOT = os.path.abspath(os.path.join(_SCRIPT_DIR, '../..'))
if _REPO_ROOT not in sys.path:
    sys.path.insert(0, _REPO_ROOT)

from statvar_imports.australia_statistics import process


class ProcessTest(unittest.TestCase):

    def test_parse_sdmx_labour_force_csv_preserves_raw_unscaled_counts(self):
        sample_csv = (
            'DATAFLOW,MEASURE,SEX,AGE,TSEST,REGION,FREQ,TIME_PERIOD,OBS_VALUE\n'
            'ABS:LF(1.0.0),M3,3,1599,10,AUS,M,2026-09,14925.782419\n'
            'ABS:LF(1.0.0),M6,1,1599,10,1,M,2026-09,105.5\n'
            'ABS:LF(1.0.0),M13,2,1599,10,2,M,2026-09,4.249\n')
        rows = process.parse_sdmx_labour_force_csv(sample_csv)
        self.assertEqual(len(rows), 3)
        self.assertEqual(
            rows[0],
            {
                'observationDate': '2026-09',
                'indicatorCd': 'LF_M3_3',
                'regionCd': 'AUS',
                'value': '14925.782419',
                'unit': '',
                'scalingFactor': '',
            },
        )
        self.assertEqual(
            rows[1],
            {
                'observationDate': '2026-09',
                'indicatorCd': 'LF_M6_1',
                'regionCd': '1',
                'value': '105.5',
                'unit': '',
                'scalingFactor': '',
            },
        )
        self.assertEqual(
            rows[2],
            {
                'observationDate': '2026-09',
                'indicatorCd': 'LF_M13_2',
                'regionCd': '2',
                'value': '4.2',
                'unit': '',
                'scalingFactor': '',
            },
        )

    def test_parse_sdmx_erp_q_csv_filters_q2_and_skips_2021(self):
        sample_csv = (
            'DATAFLOW,MEASURE,SEX,AGE,REGION,FREQ,TIME_PERIOD,OBS_VALUE\n'
            'ABS:ERP_Q(1.0.0),1,3,TOT,AUS,Q,2021-Q2,25685412\n'
            'ABS:ERP_Q(1.0.0),1,3,TOT,AUS,Q,2025-Q1,27500000\n'
            'ABS:ERP_Q(1.0.0),1,3,TOT,AUS,Q,2025-Q2,27614411\n'
            'ABS:ERP_Q(1.0.0),1,1,A04,1,Q,2025-Q2,248100\n')
        rows = process.parse_sdmx_erp_q_csv(sample_csv)
        self.assertEqual(len(rows), 2)
        self.assertEqual(
            rows[0],
            {
                'observationDate': '2025',
                'indicatorCd': 'ERP_3_TOT',
                'regionCd': 'AUS',
                'value': '27614411',
                'unit': '',
                'scalingFactor': '',
            },
        )
        self.assertEqual(
            rows[1],
            {
                'observationDate': '2025',
                'indicatorCd': 'ERP_1_A04',
                'regionCd': '1',
                'value': '248100',
                'unit': '',
                'scalingFactor': '',
            },
        )

    def test_parse_household_income_xlsx(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'Table 1.1'
        ws.cell(5, 3, '2017–18')
        ws.cell(5, 4, '2019–20')
        ws.cell(8, 1, 'Lowest quintile')
        ws.cell(8, 2, '$')
        ws.cell(8, 3, 411)
        ws.cell(8, 4, 415)
        ws.cell(21, 1, '50th (P50)')
        ws.cell(21, 2, '$')
        ws.cell(21, 3, 926)
        ws.cell(21, 4, 959)
        ws.cell(27, 1, 'Income share')
        ws.cell(28, 1, 'Lowest quintile')
        ws.cell(28, 3, 7.5)

        buf = io.BytesIO()
        wb.save(buf)
        rows = process.parse_household_income_xlsx(buf.getvalue())
        self.assertEqual(len(rows), 4)
        self.assertEqual(
            rows[0],
            {
                'observationDate': '2018',
                'indicatorCd': 'INC_Lowest quintile',
                'regionCd': 'AUS',
                'value': '411',
                'unit': 'AUD',
                'scalingFactor': '',
            },
        )
        self.assertEqual(
            rows[3],
            {
                'observationDate': '2020',
                'indicatorCd': 'INC_50th (P50)',
                'regionCd': 'AUS',
                'value': '959',
                'unit': 'AUD',
                'scalingFactor': '',
            },
        )

    def test_expand_input_files_sorts_largest_first_and_skips_empty(self):
        with tempfile.TemporaryDirectory() as tmp_dir:
            empty_file = os.path.join(tmp_dir, '0_empty.csv')
            small_file = os.path.join(tmp_dir, 'a_small.csv')
            large_file = os.path.join(tmp_dir, 'b_large.csv')
            with open(empty_file, 'w', encoding='utf-8') as f:
                f.write('')
            with open(small_file, 'w', encoding='utf-8') as f:
                f.write('a,b\n1,2\n')
            with open(large_file, 'w', encoding='utf-8') as f:
                f.write('a,b\n' + '1,2\n' * 50)

            pattern = os.path.join(tmp_dir, '*.csv')
            expanded = process._expand_input_files(pattern)
            self.assertEqual(expanded, [large_file, small_file])

    def test_merge_csv_files_keeps_single_header_and_deduplicates(self):
        with tempfile.TemporaryDirectory() as tmp_dir:
            p1 = os.path.join(tmp_dir, 'shard_00')
            p2 = os.path.join(tmp_dir, 'shard_01')
            with open(f'{p1}.csv', 'w', encoding='utf-8') as f:
                f.write(
                    'observationAbout,observationDate,value,variableMeasured,unit,scalingFactor\n'
                    'country/AUS,2021,25422788,dcid:Count_Person,,\n')
            with open(f'{p2}.csv', 'w', encoding='utf-8') as f:
                f.write(
                    'observationAbout,observationDate,value,variableMeasured,unit,scalingFactor\n'
                    'country/AUS,2021,25422788,dcid:Count_Person,,\n'
                    'country/AUS,2020,959,dcid:Weekly_Median_Income_Household,AUD,\n'
                )

            final_csv = os.path.join(tmp_dir, 'merged.csv')
            total_rows = process._merge_csv_files([p1, p2], final_csv)
            self.assertEqual(total_rows, 2)
            with open(final_csv, 'r', encoding='utf-8') as f:
                lines = f.read().splitlines()
            self.assertEqual(
                lines,
                [
                    'observationAbout,observationDate,value,variableMeasured,unit,scalingFactor',
                    'country/AUS,2021,25422788,dcid:Count_Person,,',
                    'country/AUS,2020,959,dcid:Weekly_Median_Income_Household,AUD,',
                ],
            )

    def test_merge_counters_aggregates_by_mode(self):
        with tempfile.TemporaryDirectory() as tmp_dir:
            c1 = os.path.join(tmp_dir, 'c1.csv')
            c2 = os.path.join(tmp_dir, 'c2.csv')
            with open(c1, 'w', encoding='utf-8') as f:
                f.write('key,value\n'
                        'svobs-added,100\n'
                        'existing-statvars-from-mcf,141928\n'
                        'process_start_time,1000\n'
                        'process_end_time,1050\n'
                        'processing_rate,10.0\n')
            with open(c2, 'w', encoding='utf-8') as f:
                f.write('key,value\n'
                        'svobs-added,250\n'
                        'existing-statvars-from-mcf,141928\n'
                        'process_start_time,1010\n'
                        'process_end_time,1080\n'
                        'processing_rate,20.0\n')

            final_counters = os.path.join(tmp_dir, 'merged_counters.csv')
            process._merge_counters([c1, c2], final_counters)
            with open(final_counters, 'r', encoding='utf-8') as f:
                lines = f.read().splitlines()
            self.assertEqual(
                lines,
                [
                    'key,value',
                    'existing-statvars-from-mcf,141928',
                    'process_end_time,1080',
                    'process_start_time,1000',
                    'processing_rate,15.00',
                    'svobs-added,350',
                ],
            )

    def test_merge_mcf_files_deduplicates_and_skips_empty(self):
        with tempfile.TemporaryDirectory() as tmp_dir:
            p1 = os.path.join(tmp_dir, 'shard_00')
            p2 = os.path.join(tmp_dir, 'shard_01')
            final_mcf = os.path.join(tmp_dir, 'merged_stat_vars.mcf')

            process._merge_mcf_files([p1, p2], '_stat_vars.mcf', final_mcf)
            self.assertFalse(os.path.exists(final_mcf))

            node_block = ('# Auto generated comment\n\n'
                          'Node: dcid:Count_Person\n'
                          'typeOf: dcid:StatisticalVariable\n'
                          'populationType: dcid:Person\n')
            with open(f'{p1}_stat_vars.mcf', 'w', encoding='utf-8') as f:
                f.write(node_block)
            with open(f'{p2}_stat_vars.mcf', 'w', encoding='utf-8') as f:
                f.write(node_block)

            process._merge_mcf_files([p1, p2], '_stat_vars.mcf', final_mcf)
            self.assertTrue(os.path.exists(final_mcf))
            with open(final_mcf, 'r', encoding='utf-8') as f:
                content = f.read()
            self.assertEqual(content.count('Node: dcid:Count_Person'), 1)

    @mock.patch.object(process.subprocess, 'run')
    def test_run_shard_raises_on_nonzero_exit(self, mock_run):
        mock_proc = mock.MagicMock()
        mock_proc.returncode = 1
        mock_proc.stderr = 'Processor failure details'
        mock_run.return_value = mock_proc

        with tempfile.TemporaryDirectory() as tmp_dir:
            out_prefix = os.path.join(tmp_dir, 'out_prefix')
            with self.assertRaises(RuntimeError):
                process._run_shard(
                    shard_idx=0,
                    input_file='input.csv',
                    shard_output_prefix=out_prefix,
                    shard_counters_path='counters.csv',
                    pv_map='pv.csv',
                    config_file='meta.csv',
                    existing_statvar_mcf='stat_vars.mcf',
                )


if __name__ == '__main__':
    unittest.main()
