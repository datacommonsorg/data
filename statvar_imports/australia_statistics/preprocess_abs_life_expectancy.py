# Copyright 2026 Google LLC
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#      http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.
"""Downloads and merges ABS Life Tables into life_expectancy intermediate CSVs."""

import csv
import datetime
import io
import os
import re
from typing import Dict, List, Optional, Tuple

from absl import app
from absl import flags
from absl import logging
from bs4 import BeautifulSoup
import openpyxl
import requests

from preprocess_abs_labour_force import create_retry_session

FLAGS = flags.FLAGS

flags.DEFINE_string(
    'age_csv',
    os.path.join(
        os.path.dirname(os.path.abspath(__file__)),
        'intermediate_files',
        'life_expectancy_age_aus.csv',
    ),
    'Path to intermediate life_expectancy_age_aus.csv to update.',
)
flags.DEFINE_string(
    'states_csv',
    os.path.join(
        os.path.dirname(os.path.abspath(__file__)),
        'intermediate_files',
        'life_expectancy_aus_states.csv',
    ),
    'Path to intermediate life_expectancy_aus_states.csv to update.',
)

LIFE_EXPECTANCY_LATEST_URL = (
    'https://www.abs.gov.au/statistics/people/population/life-expectancy/'
    'latest-release')
LIFE_EXPECTANCY_BASE_URL = (
    'https://www.abs.gov.au/statistics/people/population/life-expectancy')

BASELINE_LAST_DO001_START_YEAR = 2018

DO002_PLACE_LABELS: Dict[str, str] = {
    'Total New South Wales': 'wikidataId/Q3224',
    'Total Victoria': 'wikidataId/Q36687',
    'Total Queensland': 'wikidataId/Q36074',
    'Total South Australia': 'wikidataId/Q35715',
    'Total Western Australia': 'wikidataId/Q3206',
    'Total Tasmania': 'wikidataId/Q34366',
    'Total Northern Territory': 'wikidataId/Q3235',
    'Total Australian Capital Territory': 'wikidataId/Q3258',
    'Australia': 'country/AUS',
}


def resolve_latest_life_expectancy_period(session: requests.Session) -> str:
    """Resolves the latest 3-year period slug from ABS."""
    resp = session.get(LIFE_EXPECTANCY_LATEST_URL, timeout=60)
    resp.raise_for_status()
    match = re.search(r'/life-expectancy/(\d{4}-\d{4})', resp.url)
    if match:
        return match.group(1)
    soup = BeautifulSoup(resp.text, 'html.parser')
    for link in soup.find_all('a', href=True):
        m = re.search(r'/life-expectancy/(\d{4}-\d{4})/3302055001DO001',
                      link['href'])
        if m:
            return m.group(1)
    current_year = datetime.date.today().year
    for end_year in range(current_year, BASELINE_LAST_DO001_START_YEAR + 1, -1):
        period = f'{end_year - 2}-{end_year}'
        slug = period.replace('-', '')
        probe_url = (f'{LIFE_EXPECTANCY_BASE_URL}/{period}/'
                     f'3302055001DO001_{slug}.xlsx')
        probe = session.head(probe_url, timeout=30)
        if probe.status_code == 200:
            return period
    raise RuntimeError(
        f'Unable to resolve ABS Life Expectancy period from {LIFE_EXPECTANCY_LATEST_URL}'
    )


def build_do001_periods(
    latest_period: str,
    start_year: int = BASELINE_LAST_DO001_START_YEAR,
) -> List[str]:
    """Builds consecutive 3-year period slugs up to latest_period."""
    match = re.match(r'^(\d{4})-(\d{4})$', latest_period.strip())
    if not match:
        raise ValueError(f'Invalid 3-year period format: {latest_period}')
    latest_start = int(match.group(1))
    return [
        f'{year}-{year + 2}' for year in range(start_year, latest_start + 1)
    ]


def ensure_pvmap_periods(pvmap_path: str, periods: List[str]) -> None:
    """Ensures every 3-year period (YYYY-YYYY) is mapped in the PVMap file."""
    if not os.path.exists(pvmap_path):
        return
    with open(pvmap_path, 'r', encoding='utf-8') as f:
        content = f.read()

    missing_entries: List[str] = []
    for period in periods:
        if f'"{period}"' not in content and f"'{period}'" not in content:
            end_year = period.split('-')[-1]
            entry = (f'   "{period}": {{\n'
                     f'       "observationDate": "{end_year}",\n'
                     f'       "observationPeriod": "P3Y",\n'
                     f'       "value": "{{Number}}",\n'
                     f'   }},\n')
            missing_entries.append(entry)

    if not missing_entries:
        return

    closing_idx = content.rfind('}')
    if closing_idx == -1:
        raise ValueError(f'Invalid PVMap file structure in {pvmap_path}')
    updated = content[:closing_idx] + ''.join(
        missing_entries) + content[closing_idx:]
    with open(pvmap_path, 'w', encoding='utf-8') as f:
        f.write(updated)


def parse_do001_national_age_table(
        workbook_bytes: bytes) -> Dict[Tuple[str, str], str]:
    """Extracts (gender, Age) -> ex from Table 9 of 3302055001DO001."""
    wb = openpyxl.load_workbook(io.BytesIO(workbook_bytes),
                                data_only=True,
                                read_only=True)
    target_sheet = None
    for candidate in ('Table 9', 'Table 1.9', 'Table_1.9'):
        if candidate in wb.sheetnames:
            target_sheet = candidate
            break
    if not target_sheet:
        for sname in wb.sheetnames:
            if '9' in sname:
                target_sheet = sname
                break
    if not target_sheet:
        raise ValueError(
            f'Missing Table 9 in 3302055001DO001 workbook: {wb.sheetnames}')
    ws = wb[target_sheet]
    result: Dict[Tuple[str, str], str] = {}
    for row in ws.iter_rows(values_only=True):
        if not row or row[0] is None:
            continue
        age_str = str(row[0]).strip()
        if not age_str.isdigit():
            continue
        age_int = int(age_str)
        if 1 <= age_int <= 99:
            male_ex = row[4] if len(row) > 4 else None
            female_ex = row[8] if len(row) > 8 else None
            if male_ex is not None and male_ex != '':
                result[('Male', str(age_int))] = f'{float(male_ex):.2f}'
            if female_ex is not None and female_ex != '':
                result[('Female', str(age_int))] = f'{float(female_ex):.2f}'
    return result


def parse_do002_states_table(
        workbook_bytes: bytes
) -> Tuple[List[str], Dict[Tuple[str, str, str], str]]:
    """Extracts (gender, place_dcid, period) -> ex from Table 1 of 3302055001DO002."""
    wb = openpyxl.load_workbook(io.BytesIO(workbook_bytes),
                                data_only=True,
                                read_only=True)
    if 'Table 1' not in wb.sheetnames:
        raise ValueError('Missing Table 1 in 3302055001DO002 workbook')
    ws = wb['Table 1']
    rows = list(ws.iter_rows(values_only=True))
    if len(rows) < 10:
        raise ValueError('Insufficient rows in Table 1 of 3302055001DO002')

    period_row = rows[5]
    period_cols: List[Tuple[str, int, int]] = []
    periods: List[str] = []
    for col_idx, cell in enumerate(period_row):
        if not cell:
            continue
        normalized = str(cell).strip().replace('–', '-').replace('—', '-')
        if re.match(r'^\d{4}-\d{4}$', normalized):
            periods.append(normalized)
            period_cols.append((normalized, col_idx, col_idx + 1))

    values: Dict[Tuple[str, str, str], str] = {}
    for row in rows[10:]:
        if not row or len(row) < 3 or row[1] is None:
            continue
        label = str(row[1]).strip()
        place_dcid = DO002_PLACE_LABELS.get(label)
        if not place_dcid:
            continue
        for period, male_col, female_col in period_cols:
            male_val = row[male_col] if male_col < len(row) else None
            female_val = row[female_col] if female_col < len(row) else None
            if male_val is not None and male_val != '':
                val_str = f'{round(float(male_val), 1):.1f}'.rstrip('0').rstrip(
                    '.')
                values[('Male', place_dcid, period)] = val_str
            if female_val is not None and female_val != '':
                val_str = f'{round(float(female_val), 1):.1f}'.rstrip(
                    '0').rstrip('.')
                values[('Female', place_dcid, period)] = val_str
    return periods, values


def refresh_life_expectancy_csvs(
    age_csv_path: str,
    states_csv_path: str,
    session: Optional[requests.Session] = None,
) -> Tuple[int, int]:
    """Downloads ABS Life Tables and updates age and state intermediate CSVs."""
    own_session = session is None
    if own_session:
        session = create_retry_session()
    try:
        latest_period = resolve_latest_life_expectancy_period(session)
        periods_to_fetch = build_do001_periods(latest_period)

        # 1. Update life_expectancy_age_aus.csv from DO001 workbooks
        with open(age_csv_path, 'r', encoding='utf-8') as f:
            reader = csv.DictReader(f)
            age_fieldnames = list(reader.fieldnames or [])
            age_rows = list(reader)

        for period in periods_to_fetch:
            slug = period.replace('-', '')
            url = (f'{LIFE_EXPECTANCY_BASE_URL}/{period}/'
                   f'3302055001DO001_{slug}.xlsx')
            logging.info('Downloading ABS Life Tables DO001 for %s: %s', period,
                         url)
            resp = session.get(url, timeout=60)
            resp.raise_for_status()
            age_map = parse_do001_national_age_table(resp.content)
            if period not in age_fieldnames:
                age_fieldnames.append(period)
            for row in age_rows:
                key = (row['gender'].strip(), row['Age'].strip())
                if key in age_map:
                    row[period] = age_map[key]

        with open(age_csv_path, 'w', encoding='utf-8', newline='') as f:
            writer = csv.DictWriter(f, fieldnames=age_fieldnames)
            writer.writeheader()
            writer.writerows(age_rows)

        age_pvmap = os.path.join(os.path.dirname(age_csv_path),
                                 'life_expectancy_age_aus_map.py')
        ensure_pvmap_periods(age_pvmap, periods_to_fetch)

        # 2. Update life_expectancy_aus_states.csv from latest DO002 workbook
        latest_slug = latest_period.replace('-', '')
        do002_url = (f'{LIFE_EXPECTANCY_BASE_URL}/{latest_period}/'
                     f'3302055001DO002_{latest_slug}.xlsx')
        logging.info('Downloading ABS Life Tables DO002: %s', do002_url)
        resp2 = session.get(do002_url, timeout=60)
        resp2.raise_for_status()
        do002_periods, state_values = parse_do002_states_table(resp2.content)

        with open(states_csv_path, 'r', encoding='utf-8') as f:
            reader = csv.DictReader(f)
            state_fieldnames = list(reader.fieldnames or [])
            state_rows = list(reader)

        for period in do002_periods:
            if period not in state_fieldnames:
                state_fieldnames.append(period)
            for row in state_rows:
                key = (
                    row['gender'].strip(),
                    row['observationAbout'].strip(),
                    period,
                )
                if key in state_values:
                    row[period] = state_values[key]

        with open(states_csv_path, 'w', encoding='utf-8', newline='') as f:
            writer = csv.DictWriter(f, fieldnames=state_fieldnames)
            writer.writeheader()
            writer.writerows(state_rows)

        states_pvmap = os.path.join(os.path.dirname(states_csv_path),
                                    'life_expectancy_aus_states_map.py')
        ensure_pvmap_periods(states_pvmap, do002_periods)

        return len(age_rows), len(state_rows)
    finally:
        if own_session and session is not None:
            session.close()


def main(_):
    refresh_life_expectancy_csvs(FLAGS.age_csv, FLAGS.states_csv)


if __name__ == '__main__':
    app.run(main)
