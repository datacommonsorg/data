# Copyright 2026 Google LLC
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#      http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.
"""Unit tests for process_mmg."""

import os
import tempfile
import unittest

import openpyxl

from statvar_imports.feeding_america.map_meal_gap import process_mmg

_SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
_TEST_DATA_DIR = os.path.join(_SCRIPT_DIR, 'test_data')
_TMCF_PATH = os.path.join(_SCRIPT_DIR, 'FeedingAmerica_MMG.tmcf')


class ProcessMmgTest(unittest.TestCase):

    def test_format_fips_dcid(self):
        self.assertEqual(process_mmg.format_fips_dcid(1, is_state=True),
                         'geoId/01')
        self.assertEqual(process_mmg.format_fips_dcid('06', is_state=True),
                         'geoId/06')
        self.assertEqual(process_mmg.format_fips_dcid(1001, is_state=False),
                         'geoId/01001')
        self.assertEqual(process_mmg.format_fips_dcid('56045', is_state=False),
                         'geoId/56045')
        self.assertEqual(
            process_mmg.format_fips_dcid(101, geo_type='cong_dist'),
            'geoId/0101')
        self.assertEqual(
            process_mmg.format_fips_dcid('5600', geo_type='cong_dist'),
            'geoId/5600')
        self.assertEqual(process_mmg.format_fips_dcid('geoId/01'), 'geoId/01')
        self.assertEqual(process_mmg.format_fips_dcid(None), '')
        self.assertEqual(process_mmg.format_fips_dcid('n/a'), '')
        self.assertEqual(process_mmg.format_fips_dcid(-5), '')

    def test_parse_partner_row_derivations_and_na_threshold(self):
        # Case 1: State with 3 distinct thresholds (SNAP + Other Nutrition)
        row_al = {
            'fips': 1,
            'year': 2019,
            '# of food insecure persons overall': '788,250',
            '% fi ≤ low threshold': 0.5029876308277831,
            '% fi btwn thresholds': 0.15274341896606405,
            '% fi > high threshold': 0.34397716460513794,
            '% fi > low threshold': 0.4970123691722169,
            '# of food insecure children': '226,310',
            '% food insecure children in hh w/ hh incomes below 185 fpl': 0.78,
            '% food insecure children in hh w/ hh incomes above 185 fpl': 0.22,
            'cost per meal': '$3.01',
            'weighted annual food budget shortfall': '$405,611,000',
        }
        parsed_al = process_mmg.parse_partner_row(row_al, is_state=True)
        self.assertIsNotNone(parsed_al)
        self.assertEqual(parsed_al['observationAbout'], 'geoId/01')
        self.assertEqual(parsed_al['observationDate'], '2019')
        self.assertEqual(parsed_al['Count_Person_FoodInsecure'], '788250')
        self.assertEqual(
            parsed_al['Count_Person_FoodInsecure_EligibleForAllFederalBenefits'],
            '396480.0',
        )
        self.assertEqual(
            parsed_al[
                'Count_Person_FoodInsecure_EligibleForSomeFederalBenefits'],
            '120400.0',
        )
        self.assertEqual(
            parsed_al['Count_Person_FoodInsecure_IneligibleForFederalBenefits'],
            '271140.0',
        )
        self.assertEqual(parsed_al['Count_Child_FoodInsecure'], '226310')
        self.assertEqual(
            parsed_al[
                'Count_Child_FoodInsecure_EligibleForManyChildNutritionBenefits'
            ],
            '176521.8',
        )
        self.assertEqual(
            parsed_al[
                'Count_Child_FoodInsecure_IneligibleForChildNutritionBenefits'],
            '49788.2',
        )
        self.assertEqual(parsed_al['Mean_MealCost_Person_FoodSecure'], '3.01')
        self.assertEqual(parsed_al['Annual_FoodBudgetShortfall'], '405611000')

        # Case 2: State/year where '% FI Btwn Thresholds' and '% FI > High
        # Threshold' are 'n/a' or blank, falling back to '% FI > Low Threshold'
        row_az = {
            'fips': '4',
            'year': '2021',
            '# of food insecure persons overall': '751,690',
            '% fi ≤ low threshold': '48.6%',
            '% fi btwn thresholds': 'n/a',
            '% fi > high threshold': '',
            '% fi > low threshold': '51.4%',
            '# of food insecure children': '226,080',
            '% food insecure children in hh w/ hh incomes below 185 fpl': '68%',
            '% food insecure children in hh w/ hh incomes above 185 fpl': '32%',
            'cost per meal': '$3.30',
            'weighted annual food budget shortfall': '$438,117,000',
        }
        parsed_az = process_mmg.parse_partner_row(row_az, is_state=True)
        self.assertIsNotNone(parsed_az)
        self.assertEqual(parsed_az['observationAbout'], 'geoId/04')
        self.assertEqual(
            parsed_az[
                'Count_Person_FoodInsecure_EligibleForSomeFederalBenefits'],
            '',
        )
        self.assertEqual(
            parsed_az['Count_Person_FoodInsecure_IneligibleForFederalBenefits'],
            '386368.66',
        )

    def test_process_data_csv_fixtures_match_expected(self):
        with tempfile.TemporaryDirectory() as tmp_out:
            output_csv = process_mmg.process_data(
                input_dir=_TEST_DATA_DIR,
                output_dir=tmp_out,
                tmcf_path=_TMCF_PATH,
            )
            expected_csv = os.path.join(_TEST_DATA_DIR, 'expected_output.csv')
            with open(output_csv, encoding='utf-8') as actual_f, open(
                    expected_csv, encoding='utf-8') as expected_f:
                self.assertEqual(actual_f.read(), expected_f.read())

            output_tmcf = os.path.join(tmp_out, 'FeedingAmerica_MMG.tmcf')
            with open(output_tmcf, encoding='utf-8') as actual_tmcf, open(
                    _TMCF_PATH, encoding='utf-8') as expected_tmcf:
                self.assertEqual(actual_tmcf.read(), expected_tmcf.read())

    def test_process_data_excel_workbook(self):
        with tempfile.TemporaryDirectory() as tmp_in, \
                tempfile.TemporaryDirectory() as tmp_out:
            wb = openpyxl.Workbook()
            ws_state = wb.active
            ws_state.title = 'State'
            ws_state.append([
                'FIPS',
                'Year',
                '# of Food Insecure Persons Overall ',
                '% FI ≤ Low Threshold',
                '% FI Btwn Thresholds',
                '% FI > High Threshold',
                '% FI > Low Threshold',
                '# of Food Insecure Children ',
                '% food insecure Children in HH w/HH Incomes Below 185 FPL',
                '% food insecure Children in HH w/HH Incomes Above 185 FPL',
                'Cost Per Meal ',
                'Weighted Annual Food Budget Shortfall',
            ])
            ws_state.append([
                1,
                2019,
                788250,
                0.5029876308277831,
                0.15274341896606405,
                0.34397716460513794,
                0.4970123691722169,
                226310,
                0.78,
                0.22,
                3.01,
                405611000,
            ])

            ws_county = wb.create_sheet('County')
            ws_county.append([
                'FIPS',
                'Year',
                '# of Food Insecure Persons Overall',
                '% FI ≤ Low Threshold',
                '% FI Btwn Thresholds',
                '% FI > High Threshold',
                '% FI > Low Threshold',
                '# of Food Insecure Children',
                '% food insecure children in HH w/ HH incomes below 185 FPL',
                '% food insecure children in HH w/ HH incomes above 185 FPL',
                'Cost Per Meal',
                'Weighted Annual Food Budget Shortfall',
            ])
            ws_county.append([
                1001,
                2021,
                7770,
                0.483,
                None,
                None,
                0.517,
                2000,
                0.68,
                0.32,
                3.6,
                4947000,
            ])

            ws_cd = wb.create_sheet('Cong District')
            ws_cd.append([
                'FIPS',
                'Year',
                '# of Food Insecure Persons Overall',
                '% FI ≤ Low Threshold',
                '% FI Btwn Thresholds',
                '% FI > High Threshold',
                '% FI > Low Threshold',
                '# of Food Insecure Children',
                '% food insecure children in HH w/ HH incomes below 185 FPL',
                '% food insecure children in HH w/ HH incomes above 185 FPL',
            ])
            ws_cd.append([
                101,
                2021,
                108230,
                0.397,
                None,
                None,
                0.603,
                30560,
                0.77,
                0.23,
            ])

            xlsx_path = os.path.join(tmp_in, 'MMG2023_Test.xlsx')
            wb.save(xlsx_path)

            output_csv = process_mmg.process_data(
                input_dir=tmp_in,
                output_dir=tmp_out,
                tmcf_path=_TMCF_PATH,
            )
            with open(output_csv, encoding='utf-8') as f:
                lines = [line.strip() for line in f if line.strip()]
            self.assertEqual(len(lines), 4)
            self.assertEqual(
                lines[1],
                'geoId/01,2019,788250,396480.0,120400.0,271140.0,226310,176521.8,49788.2,3.01,405611000',
            )
            self.assertEqual(
                lines[2],
                'geoId/0101,2021,108230,42967.31,,65262.69,30560,23531.2,7028.8,,',
            )
            self.assertEqual(
                lines[3],
                'geoId/01001,2021,7770,3752.91,,4017.09,2000,1360.0,640.0,3.6,4947000',
            )

    def test_process_data_errors_on_invalid_inputs(self):
        with tempfile.TemporaryDirectory() as tmp_empty, \
                tempfile.TemporaryDirectory() as tmp_out:
            with self.assertRaises(RuntimeError):
                process_mmg.process_data(
                    input_dir=tmp_empty,
                    output_dir=tmp_out,
                    tmcf_path=_TMCF_PATH,
                )
            with self.assertRaises(FileNotFoundError):
                process_mmg.process_data(
                    input_dir='/nonexistent/path/for/mmg',
                    output_dir=tmp_out,
                    tmcf_path=_TMCF_PATH,
                )


if __name__ == '__main__':
    unittest.main()
