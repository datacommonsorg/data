#!/usr/bin/env python3
# Copyright 2026 Google LLC
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#      http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.
"""Unit tests for Sri Lanka Census download.py."""

import io
import os
import tempfile
import unittest
from unittest import mock

import download
import openpyxl


class DownloadTest(unittest.TestCase):

    def _create_sample_xlsx_bytes(self) -> bytes:
        wb = openpyxl.Workbook()
        # Remove default sheet
        wb.remove(wb.active)

        headers_adm0 = ['year', 'ISO3', 'ADM0_EN', 'ADM0_PCODE'
                       ] + (download.POPULATION_COLUMNS)
        ws0 = wb.create_sheet('lka_admpop_adm0_2022')
        ws0.append(headers_adm0)
        ws0.append(['2022', 'LKA', 'Sri Lanka', 'LK'] + [100] * 54)

        headers_adm1 = [
            'year',
            'ISO3',
            'ADM0_EN',
            'ADM0_PCODE',
            'ADM1_EN',
            'ADM1_PCODE',
        ] + download.POPULATION_COLUMNS
        ws1 = wb.create_sheet('lka_admpop_adm1_2022')
        ws1.append(headers_adm1)
        ws1.append(['2022', 'LKA', 'Sri Lanka', 'LK', 'Western', 'LK1'] +
                   [50] * 54)

        headers_adm2 = [
            'year',
            'ISO3',
            'ADM0_EN',
            'ADM0_PCODE',
            'ADM1_EN',
            'ADM1_PCODE',
            'ADM2_EN',
            'ADM2_PCODE',
        ] + download.POPULATION_COLUMNS
        ws2 = wb.create_sheet('lka_admpop_adm2_2022')
        ws2.append(headers_adm2)
        ws2.append([
            '2022',
            'LKA',
            'Sri Lanka',
            'LK',
            'Western',
            'LK1',
            'Colombo',
            'LK11',
        ] + [25] * 54)
        # Trailing empty row in Excel
        ws2.append([None] * len(headers_adm2))

        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()

    def test_parse_population_xlsx(self):
        xlsx_bytes = self._create_sample_xlsx_bytes()
        rows = download.parse_population_xlsx(xlsx_bytes)
        self.assertEqual(len(rows), 3)
        self.assertEqual(rows[0]['place_code'], 'LKA')
        self.assertEqual(rows[0]['year'], '2022')
        self.assertEqual(rows[0]['F_TL'], '100')
        self.assertEqual(rows[1]['place_code'], 'LK1')
        self.assertEqual(rows[2]['place_code'], 'LK11')

    def test_parse_population_csv(self):
        headers = ['\ufeffyear', 'ISO3', 'ADM1_PCODE', 'ADM2_PCODE'
                  ] + (download.POPULATION_COLUMNS)
        row = ['2023', 'LKA', 'LK1', 'LK12'] + ['200'] * 54
        csv_text = ','.join(headers) + '\n' + ','.join(row) + '\n'
        rows = download.parse_population_csv(csv_text)
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]['year'], '2023')
        self.assertEqual(rows[0]['place_code'], 'LK12')
        self.assertEqual(rows[0]['T_TL'], '200')

    def test_parse_wdi_csv_filters_hxl_and_unmapped(self):
        csv_text = (
            'Country Name,Country ISO3,Year,Indicator Name,Indicator'
            ' Code,Value\n#country+name,#country+code,#date+year,#indicator+name,#indicator+code,#indicator+value+num\nSri'
            ' Lanka,LKA,2022,Arable land,AG.LND.ARBL.HA,1300000\nSri'
            ' Lanka,LKA,2022,Unmapped,UNMAPPED.CODE,999\n')
        rows = download.parse_wdi_csv(csv_text,
                                      allowed_indicators={'AG.LND.ARBL.HA'})
        self.assertEqual(
            rows,
            [{
                'Country ISO3': 'LKA',
                'Year': '2022',
                'Indicator Code': 'AG.LND.ARBL.HA',
                'Value': '1300000',
            }],
        )

    @mock.patch('download.download_util.request_url')
    def test_download_and_process_end_to_end(self, mock_request_url):
        xlsx_bytes = self._create_sample_xlsx_bytes()
        pop_csv_bytes = (
            ','.join(['year', 'ISO3'] + download.POPULATION_COLUMNS) + '\n' +
            ','.join(['2023', 'LKA'] + ['110'] * 54) + '\n').encode('utf-8')
        wdi_csv_bytes = (
            'Country Name,Country ISO3,Year,Indicator Name,Indicator'
            ' Code,Value\n#country+name,#country+code,#date+year,#indicator+name,#indicator+code,#indicator+value+num\nSri'
            ' Lanka,LKA,2022,Test,AG.LND.ARBL.HA,123\n').encode('utf-8')

        def side_effect(url, output='text', **_kwargs):
            if output == 'json':
                return {'success': True, 'result': {'resources': []}}
            if url.endswith('.xlsx'):
                return xlsx_bytes
            if 'admpop_adm' in url:
                return pop_csv_bytes
            return wdi_csv_bytes

        mock_request_url.side_effect = side_effect

        with tempfile.TemporaryDirectory() as tmp_dir:
            pv_map_path = os.path.join(tmp_dir, 'pvmap.csv')
            with open(pv_map_path, 'w', encoding='utf-8') as f:
                f.write('key,property,value\n')
                f.write('AG.LND.ARBL.HA,variableMeasured,dcid:Area_LandCover\n')

            out_dir = os.path.join(tmp_dir, 'input_files')
            download.download_and_process(out_dir, pv_map_path)

            self.assertTrue(
                os.path.exists(os.path.join(out_dir,
                                            'srilanka_population.csv')))
            self.assertTrue(
                os.path.exists(os.path.join(out_dir,
                                            'srilanka_agriculture.csv')))
            self.assertTrue(
                os.path.exists(os.path.join(out_dir, 'srilanka_education.csv')))


if __name__ == '__main__':
    unittest.main()
