#!/usr/bin/env python3
# Copyright 2026 Google LLC
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#      http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.
"""Downloads and preprocesses Sri Lanka Census datasets from HDX."""

import csv
import io
import os
import re
import sys
import tempfile
from typing import Any

_CODEDIR = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(1, os.path.join(_CODEDIR, '../../util/'))

from absl import app
from absl import flags
from absl import logging
import download_util
import openpyxl

_FLAGS = flags.FLAGS
flags.DEFINE_string(
    'output_dir',
    os.path.join(_CODEDIR, 'input_files'),
    'Directory path to save the preprocessed input CSV files.',
)
flags.DEFINE_string(
    'pv_map',
    os.path.join(_CODEDIR, 'srilanka_census_pvmap.csv'),
    'Path to the property-value map CSV file used to filter valid indicators.',
)

HDX_PACKAGE_SHOW_URL = 'https://data.humdata.org/api/3/action/package_show'

POPULATION_DATASET_ID = 'cod-ps-lka'
AGRICULTURE_DATASET_ID = (
    'world-bank-agriculture-and-rural-development-indicators-for-sri-lanka')
EDUCATION_DATASET_ID = 'world-bank-education-indicators-for-sri-lanka'

FALLBACK_POPULATION_URLS = {
    'lka_admpop_2022_v2.xlsx': (
        'https://data.humdata.org/dataset/e982aec3-478e-4910-ae5c-072ae2cf6933/'
        'resource/5e43e0da-675e-4efd-978c-3fa9a10c724e/download/'
        'lka_admpop_2022_v2.xlsx'),
    'lka_admpop_adm0_2023.csv': (
        'https://data.humdata.org/dataset/e982aec3-478e-4910-ae5c-072ae2cf6933/'
        'resource/c309ff4c-f435-4cad-85ea-97c39a963f90/download/'
        'lka_admpop_adm0_2023.csv'),
    'lka_admpop_adm1_2023.csv': (
        'https://data.humdata.org/dataset/e982aec3-478e-4910-ae5c-072ae2cf6933/'
        'resource/62dbf08d-a2bb-42bd-bdb5-aa5b6903ea98/download/'
        'lka_admpop_adm1_2023.csv'),
    'lka_admpop_adm2_2023.csv': (
        'https://data.humdata.org/dataset/e982aec3-478e-4910-ae5c-072ae2cf6933/'
        'resource/4405f46a-e4b3-4778-9290-131613a7ce71/download/'
        'lka_admpop_adm2_2023.csv'),
}

FALLBACK_AGRICULTURE_URL = (
    'https://data.humdata.org/dataset/5f03f0d6-bdb0-48bd-ab57-548cb97d7a9f/'
    'resource/a8ec11a2-feea-48c9-86f2-8d2866ce4b40/download/'
    'agriculture-and-rural-development_lka.csv')

FALLBACK_EDUCATION_URL = (
    'https://data.humdata.org/dataset/37a9ce34-ba19-47cf-8b68-62f895c30c17/'
    'resource/00a79f05-f4c7-4232-b1d8-470ade19611b/download/'
    'education_lka.csv')

POPULATION_COLUMNS = [
    'F_TL',
    'M_TL',
    'T_TL',
    'F_00_04',
    'F_05_09',
    'F_10_14',
    'F_15_19',
    'F_20_24',
    'F_25_29',
    'F_30_34',
    'F_35_39',
    'F_40_44',
    'F_45_49',
    'F_50_54',
    'F_55_59',
    'F_60_64',
    'F_65_69',
    'F_70_74',
    'F_75_79',
    'F_80Plus',
    'M_00_04',
    'M_05_09',
    'M_10_14',
    'M_15_19',
    'M_20_24',
    'M_25_29',
    'M_30_34',
    'M_35_39',
    'M_40_44',
    'M_45_49',
    'M_50_54',
    'M_55_59',
    'M_60_64',
    'M_65_69',
    'M_70_74',
    'M_75_79',
    'M_80Plus',
    'T_00_04',
    'T_05_09',
    'T_10_14',
    'T_15_19',
    'T_20_24',
    'T_25_29',
    'T_30_34',
    'T_35_39',
    'T_40_44',
    'T_45_49',
    'T_50_54',
    'T_55_59',
    'T_60_64',
    'T_65_69',
    'T_70_74',
    'T_75_79',
    'T_80Plus',
]

POPULATION_XLSX_SHEETS = [
    'lka_admpop_adm0_2022',
    'lka_admpop_adm1_2022',
    'lka_admpop_adm2_2022',
]


def load_mapped_keys(pv_map_path: str) -> set[str]:
    """Loads mapped keys from the PV map CSV."""
    if not pv_map_path or not os.path.exists(pv_map_path):
        raise FileNotFoundError(f'PV map file not found: {pv_map_path}')
    mapped_keys = set()
    with open(pv_map_path, mode='r', encoding='utf-8-sig') as f:
        reader = csv.DictReader(f)
        for row in reader:
            key = (row.get('key') or '').strip()
            prop = (row.get('property') or '').strip()
            if key and prop == 'variableMeasured':
                mapped_keys.add(key)
    return mapped_keys


def get_hdx_resources(dataset_id: str) -> dict[str, str]:
    """Queries HDX CKAN API for resource download URLs in a dataset."""
    params = {'id': dataset_id}
    logging.info('GET %s params=%s', HDX_PACKAGE_SHOW_URL, params)
    try:
        resp = download_util.request_url(
            url=HDX_PACKAGE_SHOW_URL,
            params=params,
            output='json',
            retries=3,
            retry_secs=5,
        )
        if not isinstance(resp, dict) or not resp.get('success'):
            logging.warning(
                'Unexpected HDX metadata response for %s; using fallback URLs.',
                dataset_id,
            )
            return {}
        resources = resp.get('result', {}).get('resources', [])
        url_map = {}
        for res in resources:
            name = (res.get('name') or '').strip()
            url = (res.get('download_url') or res.get('url') or '').strip()
            if name and url:
                url_map[name] = url
        return url_map
    except Exception as e:  # pylint: disable=broad-except
        logging.warning(
            'Failed to fetch HDX metadata for %s (%s); using fallback URLs.',
            dataset_id,
            e,
        )
        return {}


def _extract_place_code(row: dict[str, Any]) -> str:
    """Extracts the most granular place code from a population record."""
    for col in ('ADM2_PCODE', 'ADM1_PCODE', 'ISO3', 'ADM0_PCODE'):
        val = row.get(col)
        if val is not None and str(val).strip():
            code = str(val).strip()
            return 'LKA' if code == 'LK' else code
    return ''


def _format_cell_value(val: Any) -> str:
    """Formats a cell value cleanly without trailing .0 on integer floats."""
    if val is None:
        return ''
    if isinstance(val, float) and val.is_integer():
        return str(int(val))
    return str(val).strip()


def parse_population_xlsx(xlsx_bytes: bytes) -> list[dict[str, str]]:
    """Extracts ADM0, ADM1, and ADM2 population rows from the 2022 Excel file."""
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes), data_only=True)
    records = []
    for sheet_name in POPULATION_XLSX_SHEETS:
        if sheet_name not in wb.sheetnames:
            raise ValueError(
                f'Expected sheet {sheet_name} not found in workbook'
                f' {wb.sheetnames}')
        ws = wb[sheet_name]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        headers = [
            str(h).strip().lstrip('\ufeff') if h is not None else ''
            for h in rows[0]
        ]
        for raw_row in rows[1:]:
            row_dict = dict(zip(headers, raw_row))
            year_str = _format_cell_value(row_dict.get('year'))
            place_code = _extract_place_code(row_dict)
            if not year_str or not place_code:
                continue
            out_row = {'year': year_str, 'place_code': place_code}
            for col in POPULATION_COLUMNS:
                out_row[col] = _format_cell_value(row_dict.get(col))
            records.append(out_row)
    return records


def parse_population_csv(csv_text: str) -> list[dict[str, str]]:
    """Extracts population rows from an HDX population CSV file."""
    clean_text = csv_text.lstrip('\ufeff')
    reader = csv.DictReader(io.StringIO(clean_text))
    records = []
    for row_dict in reader:
        year_str = (row_dict.get('year') or '').strip()
        if year_str.startswith('#'):
            continue
        place_code = _extract_place_code(row_dict)
        if not year_str or not place_code:
            continue
        out_row = {'year': year_str, 'place_code': place_code}
        for col in POPULATION_COLUMNS:
            out_row[col] = _format_cell_value(row_dict.get(col))
        records.append(out_row)
    return records


def parse_wdi_csv(
        csv_text: str,
        allowed_indicators: set[str] | None = None) -> list[dict[str, str]]:
    """Parses an HDX World Bank indicator CSV, dropping HXL tags and unmapped codes."""
    clean_text = csv_text.lstrip('\ufeff')
    reader = csv.DictReader(io.StringIO(clean_text))
    records = []
    seen = set()
    for row in reader:
        iso3 = (row.get('Country ISO3') or '').strip()
        year = (row.get('Year') or '').strip()
        indicator_code = (row.get('Indicator Code') or '').strip()
        value = (row.get('Value') or '').strip()
        if not iso3 or iso3.startswith('#') or not year or not indicator_code:
            continue
        if value == '':
            continue
        if allowed_indicators and indicator_code not in allowed_indicators:
            continue
        dedup_key = (iso3, year, indicator_code)
        if dedup_key in seen:
            continue
        seen.add(dedup_key)
        records.append({
            'Country ISO3': iso3,
            'Year': year,
            'Indicator Code': indicator_code,
            'Value': value,
        })
    return records


def write_csv_atomic(output_path: str, fieldnames: list[str],
                     rows: list[dict[str, str]]) -> None:
    """Atomically writes rows to a CSV file."""
    output_dir = os.path.dirname(os.path.abspath(output_path))
    os.makedirs(output_dir, exist_ok=True)
    fd, tmp_path = tempfile.mkstemp(dir=output_dir,
                                    prefix='.tmp_',
                                    suffix='.csv')
    try:
        with os.fdopen(fd, mode='w', encoding='utf-8', newline='') as f:
            writer = csv.DictWriter(f,
                                    fieldnames=fieldnames,
                                    lineterminator='\n')
            writer.writeheader()
            writer.writerows(rows)
        os.replace(tmp_path, output_path)
    except Exception:
        if os.path.exists(tmp_path):
            os.remove(tmp_path)
        raise


def download_and_process(output_dir: str, pv_map_path: str) -> None:
    """Downloads all Sri Lanka datasets from HDX and writes input CSVs."""
    os.makedirs(output_dir, exist_ok=True)
    mapped_indicators = load_mapped_keys(pv_map_path)
    logging.info(
        'Loaded %d mapped variable keys from %s',
        len(mapped_indicators),
        pv_map_path,
    )

    # 1. Population (2022 XLSX + 2023 CSVs)
    pop_resources = get_hdx_resources(POPULATION_DATASET_ID)
    for fallback_name, fallback_url in FALLBACK_POPULATION_URLS.items():
        if fallback_name not in pop_resources:
            pop_resources[fallback_name] = fallback_url

    population_rows = []
    xlsx_url = pop_resources.get('lka_admpop_2022_v2.xlsx')
    if not xlsx_url:
        raise RuntimeError('Missing URL for lka_admpop_2022_v2.xlsx')
    logging.info('Downloading 2022 population workbook from %s...', xlsx_url)
    xlsx_bytes = download_util.request_url(url=xlsx_url,
                                           output='bytes',
                                           timeout=60,
                                           retries=3,
                                           retry_secs=5)
    if not isinstance(xlsx_bytes, bytes) or not xlsx_bytes:
        raise RuntimeError(f'Failed to download bytes from {xlsx_url}')
    population_rows.extend(parse_population_xlsx(xlsx_bytes))

    csv_resource_names = sorted(
        name for name in pop_resources
        if re.match(r'^lka_admpop_adm[012]_\d{4}\.csv$', name, re.IGNORECASE))
    for res_name in csv_resource_names:
        csv_url = pop_resources[res_name]
        logging.info('Downloading %s from %s...', res_name, csv_url)
        csv_bytes = download_util.request_url(url=csv_url,
                                              output='bytes',
                                              timeout=60,
                                              retries=3,
                                              retry_secs=5)
        if not isinstance(csv_bytes, bytes) or not csv_bytes:
            raise RuntimeError(f'Failed to download bytes from {csv_url}')
        population_rows.extend(
            parse_population_csv(csv_bytes.decode('utf-8-sig')))

    if not population_rows:
        raise RuntimeError('No population rows extracted from HDX.')

    pop_output_path = os.path.join(output_dir, 'srilanka_population.csv')
    write_csv_atomic(
        pop_output_path,
        ['year', 'place_code'] + POPULATION_COLUMNS,
        population_rows,
    )
    logging.info('Wrote %d population rows to %s', len(population_rows),
                 pop_output_path)

    # 2. Agriculture and Rural Development
    agri_resources = get_hdx_resources(AGRICULTURE_DATASET_ID)
    agri_url = (agri_resources.get('agriculture-and-rural-development_lka.csv')
                or FALLBACK_AGRICULTURE_URL)
    logging.info('Downloading agriculture dataset from %s...', agri_url)
    agri_bytes = download_util.request_url(url=agri_url,
                                           output='bytes',
                                           timeout=60,
                                           retries=3,
                                           retry_secs=5)
    if not isinstance(agri_bytes, bytes) or not agri_bytes:
        raise RuntimeError(f'Failed to download bytes from {agri_url}')
    agri_rows = parse_wdi_csv(agri_bytes.decode('utf-8-sig'),
                              allowed_indicators=mapped_indicators)
    if not agri_rows:
        raise RuntimeError('No agriculture rows extracted from HDX.')
    agri_output_path = os.path.join(output_dir, 'srilanka_agriculture.csv')
    write_csv_atomic(
        agri_output_path,
        ['Country ISO3', 'Year', 'Indicator Code', 'Value'],
        agri_rows,
    )
    logging.info('Wrote %d agriculture rows to %s', len(agri_rows),
                 agri_output_path)

    # 3. Education
    edu_resources = get_hdx_resources(EDUCATION_DATASET_ID)
    edu_url = edu_resources.get('education_lka.csv') or FALLBACK_EDUCATION_URL
    logging.info('Downloading education dataset from %s...', edu_url)
    edu_bytes = download_util.request_url(url=edu_url,
                                          output='bytes',
                                          timeout=60,
                                          retries=3,
                                          retry_secs=5)
    if not isinstance(edu_bytes, bytes) or not edu_bytes:
        raise RuntimeError(f'Failed to download bytes from {edu_url}')
    edu_rows = parse_wdi_csv(edu_bytes.decode('utf-8-sig'),
                             allowed_indicators=mapped_indicators)
    if not edu_rows:
        raise RuntimeError('No education rows extracted from HDX.')
    edu_output_path = os.path.join(output_dir, 'srilanka_education.csv')
    write_csv_atomic(
        edu_output_path,
        ['Country ISO3', 'Year', 'Indicator Code', 'Value'],
        edu_rows,
    )
    logging.info('Wrote %d education rows to %s', len(edu_rows),
                 edu_output_path)


def main(_) -> None:
    try:
        download_and_process(_FLAGS.output_dir, _FLAGS.pv_map)
    except Exception as e:  # pylint: disable=broad-except
        logging.fatal(
            'Failed to download and process Sri Lanka Census datasets: %s',
            e,
            exc_info=True,
        )


if __name__ == '__main__':
    app.run(main)
