# Copyright 2026 Google LLC
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#     https://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.

"""Preprocessing script for Commerce EDA Persistent Poverty Counties dataset.

This script ingests the raw Persistent Poverty Counties dataset (downloaded
from U.S. Economic Development Administration (EDA) / Department of Commerce),
standardizes geographic identifiers (5-digit county FIPS codes across 50 US
states, DC, and Puerto Rico, and 5-digit island territory codes), validates
poverty percentage rates across 1990, 2000, 2020, and 2021, and generates the
normalized cleaned CSV for stat_var_processor.py.
"""

import os
import re
import tempfile

from absl import app, flags, logging
import openpyxl
import pandas as pd

MODULE_DIR = os.path.dirname(os.path.abspath(__file__))

DEFAULT_SOURCE_CSV = os.path.join(MODULE_DIR, "input_files", "Poverty.csv")
DEFAULT_SOURCE_XLSX = os.path.join(MODULE_DIR, "input_files", "EDA_FY23_PPCs.xlsx")
DEFAULT_RAW_CSV = os.path.join(MODULE_DIR, "output", "Poverty_original.csv")
CLEANED_CSV = os.path.join(MODULE_DIR, "output", "Poverty_cleaned.csv")

FLAGS = flags.FLAGS
flags.DEFINE_string(
    "source_path",
    None,
    "Path to source file (.csv or .xlsx). If not specified, automatically "
    "resolves from input_files/Poverty.csv, input_files/EDA_FY23_PPCs.xlsx, "
    "or output/Poverty_original.csv.",
)
flags.DEFINE_string("cleaned_csv_path", CLEANED_CSV, "Path to save cleaned output CSV.")
flags.DEFINE_integer("min_county_count", 3000, "Minimum number of valid counties expected.")

# Valid 2-digit US State and Territory FIPS codes
VALID_STATE_FIPS = {
    # 50 States + DC
    "01", "02", "04", "05", "06", "08", "09", "10", "11", "12", "13", "15",
    "16", "17", "18", "19", "20", "21", "22", "23", "24", "25", "26", "27",
    "28", "29", "30", "31", "32", "33", "34", "35", "36", "37", "38", "39",
    "40", "41", "42", "44", "45", "46", "47", "48", "49", "50", "51", "53",
    "54", "55", "56",
    # Territories: American Samoa, Guam, Northern Mariana Islands, Puerto Rico, Virgin Islands
    "60", "66", "69", "72", "78",
}

# Island territories where most recent estimate is from 2020 Decennial Census
ISLAND_TERRITORY_FIPS = {"60", "66", "69", "78"}

COLUMN_RENAME_MAP = {
    "GEOID": "GEOID",
    "1990 Decennial Census, % in Poverty": "poverty_rate_1990",
    "2000 Decennial Census, % in Poverty": "poverty_rate_2000",
    "Most Recent Estimate, % in Poverty*": "poverty_rate_recent",
    "Most Recent Estimate, % in Poverty": "poverty_rate_recent",
}


def clean_geoid(val):
    """Standardizes GEOIDs to 5 digits and validates against US state FIPS.

    Rejects state summary entries of the form XX000.
    """
    if pd.isna(val):
        return None
    s = str(val).strip()
    match = re.match(r"^(\d+)(?:\.0+)?$", s)
    if not match:
        return None
    digits = match.group(1)
    if len(digits) == 4:
        digits = digits.zfill(5)
    if len(digits) == 5 and digits[:2] in VALID_STATE_FIPS and digits[2:] != "000":
        return digits
    return None


def _extract_dataframe_from_excel(excel_path):
    """Extracts Underlying_Data sheet from an Excel workbook into a pandas DataFrame."""
    wb = openpyxl.load_workbook(excel_path, data_only=True)
    try:
        target_sheet = "Underlying_Data" if "Underlying_Data" in wb.sheetnames else wb.sheetnames[0]
        ws = wb[target_sheet]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            raise ValueError(f"Excel sheet '{target_sheet}' is empty.")

        # Find row with GEOID header
        header_idx = None
        for idx, r in enumerate(rows[:10]):
            row_str = [str(c).strip() for c in r if c is not None]
            if "GEOID" in row_str or any("1990" in c for c in row_str):
                header_idx = idx
                break

        if header_idx is None:
            header_idx = 2 if len(rows) > 2 else 0

        headers = [("" if c is None else str(c).strip()) for c in rows[header_idx]]
        data_rows = []
        for r in rows[header_idx + 1:]:
            if not any(r):
                continue
            data_rows.append([("" if c is None else str(c).strip()) for c in r[:len(headers)]])

        return pd.DataFrame(data_rows, columns=headers)
    finally:
        wb.close()


def resolve_source_file_path(requested_path=None):
    """Finds the raw source file, raising FileNotFoundError if missing."""
    if requested_path:
        if os.path.exists(requested_path) and os.path.getsize(requested_path) > 0:
            return requested_path
        raise FileNotFoundError(f"Specified source file not found or empty: {requested_path}")

    candidates = [
        DEFAULT_SOURCE_CSV,
        DEFAULT_SOURCE_XLSX,
        DEFAULT_RAW_CSV,
    ]
    for candidate in candidates:
        if os.path.exists(candidate) and os.path.getsize(candidate) > 0:
            return candidate

    raise FileNotFoundError(
        "No downloaded source file found. Please run download_poverty.py first to "
        f"fetch the dataset, or specify --source_path. Checked: {candidates}"
    )


def preprocess_poverty(src_path=DEFAULT_SOURCE_CSV, dst_path=CLEANED_CSV, min_county_count=3000):
    """Preprocesses the raw Poverty dataset into cleaned format with normalized columns."""
    logging.info("Preprocessing source Poverty dataset from %s...", src_path)
    if not os.path.exists(src_path) or os.path.getsize(src_path) == 0:
        logging.error("Source file does not exist or is empty: %s", src_path)
        raise ValueError(f"Source file does not exist or is empty: {src_path}")

    if src_path.lower().endswith((".xlsx", ".xls")):
        df = _extract_dataframe_from_excel(src_path)
    else:
        # Detect header row index by scanning first lines for 'GEOID'
        skip = 0
        with open(src_path, "r", encoding="utf-8", errors="ignore") as f:
            for idx in range(10):
                line = f.readline()
                if not line:
                    break
                if "GEOID" in line or "1990" in line:
                    skip = idx
                    break
        df = pd.read_csv(src_path, skiprows=skip, dtype=str)

    if df.empty:
        logging.error("Source dataset is empty: %s", src_path)
        raise ValueError(f"Source dataset is empty: {src_path}")

    # Strip column headers to avoid fragile whitespace issues
    df.columns = df.columns.str.strip()

    # Check if at least GEOID and the poverty columns are found
    if "GEOID" not in df.columns:
        logging.error("Missing required column 'GEOID' in source dataset")
        raise ValueError("Missing required column 'GEOID' in source dataset")

    # Rename columns to standard names
    rename_dict = {}
    for col in df.columns:
        clean_col = col.rstrip("*").strip()
        for k, v in COLUMN_RENAME_MAP.items():
            if col == k or clean_col == k.rstrip("*").strip():
                rename_dict[col] = v
                break

    # Guard against silent year corruption if Data Source column is present
    data_source_cols = [c for c in df.columns if "Data Source" in c]
    df = df.rename(columns=rename_dict)

    required_cols = [
        "GEOID", "poverty_rate_1990", "poverty_rate_2000", "poverty_rate_recent"
    ]
    missing = [c for c in required_cols if c not in df.columns]
    if missing:
        logging.error("Missing required columns in source dataset: %s", missing)
        raise ValueError(f"Missing required columns in source dataset: {missing}")

    # Standardize and validate GEOIDs
    df["GEOID"] = df["GEOID"].apply(clean_geoid)
    df = df.dropna(subset=["GEOID"])

    if data_source_cols:
        ds_col = data_source_cols[0]
        is_terr = df["GEOID"].str[:2].isin(ISLAND_TERRITORY_FIPS)
        for idx, row in df.iterrows():
            val = str(row.get(ds_col, "")).strip()
            if val and val != "nan":
                m = re.search(r"(\d{4})\s*$", val)
                if m:
                    yr = m.group(1)
                    expected_yr = "2020" if is_terr.loc[idx] else "2021"
                    if yr != expected_yr:
                        logging.error(
                            "Unexpected survey year %s in %s for GEOID %s (expected %s)",
                            yr,
                            ds_col,
                            row["GEOID"],
                            expected_yr,
                        )
                        raise ValueError(
                            f"Unexpected survey year {yr} in {ds_col} for "
                            f"GEOID {row['GEOID']} (expected {expected_yr})"
                        )

    # Coerce and validate poverty values within [0.0, 100.0]
    raw_poverty_cols = [
        "poverty_rate_1990", "poverty_rate_2000", "poverty_rate_recent"
    ]
    for col in raw_poverty_cols:
        df[col] = pd.to_numeric(df[col].astype(str).str.strip(), errors="coerce")
        invalid_mask = df[col].notna() & ((df[col] < 0.0) | (df[col] > 100.0))
        if invalid_mask.any():
            logging.warning(
                "Found %d out-of-bounds values in %s; setting to NaN",
                invalid_mask.sum(),
                col,
            )
            df.loc[invalid_mask, col] = None

    # Split recent poverty rate:
    # 2020 for Island Territories (AS, GU, MP, VI), 2021 for 50 States + DC + PR
    is_territory = df["GEOID"].str[:2].isin(ISLAND_TERRITORY_FIPS)
    df["poverty_rate_2020"] = df["poverty_rate_recent"].where(is_territory, None)
    df["poverty_rate_2021"] = df["poverty_rate_recent"].where(~is_territory, None)

    # Keep rows that have at least one valid poverty rate observation
    poverty_cols = [
        "poverty_rate_1990", "poverty_rate_2000", "poverty_rate_2020",
        "poverty_rate_2021"
    ]
    df = df.dropna(subset=poverty_cols, how="all")

    # Keep target columns only
    target_cols = ["GEOID"] + poverty_cols
    df = df[target_cols]

    # Verify sanity threshold
    if len(df) < min_county_count:
        logging.error(
            "Sanity check failed: Expected at least %d counties, but found %d.",
            min_county_count,
            len(df),
        )
        raise ValueError(
            f"Sanity check failed: Expected at least {min_county_count} "
            f"counties, but found {len(df)}."
        )

    # Atomic write to destination file
    dst_dir = os.path.dirname(os.path.abspath(dst_path))
    os.makedirs(dst_dir, exist_ok=True)
    with tempfile.NamedTemporaryFile(
        "w", dir=dst_dir, delete=False, suffix=".tmp", encoding="utf-8"
    ) as tmp_file:
        df.to_csv(tmp_file.name, index=False)
        temp_path = tmp_file.name

    os.replace(temp_path, dst_path)
    logging.info("Poverty dataset cleaned and saved successfully to %s!", dst_path)
    logging.info("Shape: %s", df.shape)
    return dst_path


def main(argv):
    """Main entrypoint for preprocessing the poverty dataset."""
    del argv  # Unused
    source_path = resolve_source_file_path(FLAGS.source_path)
    preprocess_poverty(
        src_path=source_path,
        dst_path=FLAGS.cleaned_csv_path,
        min_county_count=FLAGS.min_county_count,
    )


if __name__ == "__main__":
    app.run(main)
