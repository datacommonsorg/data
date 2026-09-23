# Copyright 2026 Google LLC
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#     https://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.

"""Download script for Commerce EDA Persistent Poverty Counties (PPC) dataset.

This script fetches the official Persistent Poverty Counties dataset from the
U.S. Economic Development Administration (EDA) / Department of Commerce:
  https://www.eda.gov/performance/tools/ (EDA_FY23_PPCs.xlsx)
It downloads the official Excel workbook (with automatic mirror failover),
supports ingesting directly from an existing input file, extracts the underlying
county-level poverty data table (3,241 places across 1990, 2000, and 2020/2021),
and stages the raw CSV and workbook under input_files/ and output/ for preprocessing.
"""

import csv
import io
import os
import shutil
import tempfile
import time
from urllib import parse
from absl import app, flags, logging
import openpyxl
import requests

MODULE_DIR = os.path.dirname(os.path.abspath(__file__))

EDA_PPC_XLSX_URL = (
    "https://www.eda.gov/sites/default/files/2023-03/EDA_FY23_PPCs.xlsx"
)
EDA_PPC_MIRROR_URL = (
    "https://web.archive.org/web/20250308204521if_/https://www.eda.gov/sites/default/files/2023-03/EDA_FY23_PPCs.xlsx"
)

DEFAULT_OUTPUT_DIR = os.path.join(MODULE_DIR, "input_files")
DEFAULT_OUTPUT_XLSX = os.path.join(DEFAULT_OUTPUT_DIR, "EDA_FY23_PPCs.xlsx")
DEFAULT_INPUT_CSV = os.path.join(DEFAULT_OUTPUT_DIR, "Poverty.csv")
DEFAULT_RAW_CSV_FILE = os.path.join(MODULE_DIR, "output", "Poverty_original.csv")

HTTP_HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (X11; Linux x86_64) AppleWebKit/537.36 "
        "(KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36"
    ),
    "Accept": "*/*",
}

FLAGS = flags.FLAGS
flags.DEFINE_string(
    "source_url",
    EDA_PPC_XLSX_URL,
    "Primary URL to download the Persistent Poverty Counties Excel workbook.",
)
flags.DEFINE_string(
    "mirror_url",
    EDA_PPC_MIRROR_URL,
    "Fallback mirror URL to download the Persistent Poverty Counties Excel workbook.",
)
flags.DEFINE_string(
    "input_file",
    None,
    "Optional path to a local input file (.xlsx or .csv) to use instead of downloading.",
)
flags.DEFINE_string(
    "output_dir",
    DEFAULT_OUTPUT_DIR,
    "Directory to store downloaded source and input files.",
)
flags.DEFINE_string(
    "output_xlsx_path",
    DEFAULT_OUTPUT_XLSX,
    "Destination path to save the downloaded source Excel workbook.",
)
flags.DEFINE_string(
    "output_csv_path",
    DEFAULT_INPUT_CSV,
    "Destination path to save the extracted input CSV file.",
)
flags.DEFINE_string(
    "raw_csv_path",
    DEFAULT_RAW_CSV_FILE,
    "Optional path to save an extracted raw CSV copy of the workbook.",
)
flags.DEFINE_integer(
    "max_retries",
    3,
    "Maximum number of download retry attempts per URL.",
)
flags.DEFINE_integer(
    "timeout",
    60,
    "HTTP request timeout in seconds.",
)


def download_file(
    download_url,
    output_path,
    session=None,
    max_retries=3,
    backoff_factor=1.5,
    timeout=60,
):
    """Downloads a file from download_url and saves it atomically to output_path."""
    logging.info("Downloading file from: %s", download_url)
    dst_dir = os.path.dirname(os.path.abspath(output_path))
    os.makedirs(dst_dir, exist_ok=True)

    req_session = session or requests.Session()
    last_err = None

    for attempt in range(1, max_retries + 1):
        try:
            logging.info(
                "Downloading %s (attempt %d/%d)...", download_url, attempt, max_retries
            )
            response = req_session.get(download_url, headers=HTTP_HEADERS, timeout=timeout)
            response.raise_for_status()
            content = response.content
            if not content:
                raise RuntimeError(f"Empty response body received from {download_url}")

            with tempfile.NamedTemporaryFile(
                "wb", dir=dst_dir, delete=False, suffix=".tmp"
            ) as tmp_file:
                tmp_file.write(content)
                temp_path = tmp_file.name

            os.replace(temp_path, output_path)
            logging.info(
                "Download completed successfully. Saved %d bytes to %s",
                len(content),
                output_path,
            )
            return content
        except Exception as e:
            last_err = e
            logging.warning(
                "Attempt %d/%d failed to download %s: %s",
                attempt,
                max_retries,
                download_url,
                e,
            )
            if attempt < max_retries:
                time.sleep(backoff_factor ** (attempt - 1))

    logging.error(
        "Download failed after %d attempts for %s: %s",
        max_retries,
        download_url,
        last_err,
    )
    raise RuntimeError(
        f"Download failed after {max_retries} attempts for {download_url}: {last_err}"
    ) from last_err


def extract_sheet_to_csv(
    excel_source,
    csv_output_path,
    target_sheet_name="Underlying_Data",
):
    """Extracts the underlying data worksheet from Excel workbook to a raw CSV."""
    if isinstance(excel_source, bytes):
        wb = openpyxl.load_workbook(io.BytesIO(excel_source), data_only=True)
    else:
        wb = openpyxl.load_workbook(excel_source, data_only=True)

    sheet_names = wb.sheetnames
    selected_sheet = None
    if target_sheet_name in sheet_names:
        selected_sheet = target_sheet_name
    else:
        for name in sheet_names:
            if any(k in name.lower() for k in ["underlying", "poverty", "data", "fy23"]):
                selected_sheet = name
                break
    if not selected_sheet:
        selected_sheet = sheet_names[0]

    logging.info("Extracting sheet '%s' from Excel workbook...", selected_sheet)
    ws = wb[selected_sheet]

    dst_dir = os.path.dirname(os.path.abspath(csv_output_path))
    os.makedirs(dst_dir, exist_ok=True)

    row_count = 0
    with tempfile.NamedTemporaryFile(
        "w", dir=dst_dir, delete=False, suffix=".tmp", encoding="utf-8", newline=""
    ) as tmp:
        writer = csv.writer(tmp)
        for row in ws.iter_rows(values_only=True):
            if not any(row):
                continue
            writer.writerow([("" if c is None else str(c)) for c in row])
            row_count += 1
        tmp_path = tmp.name

    os.replace(tmp_path, csv_output_path)
    logging.info(
        "Extracted %d rows from sheet '%s' to %s",
        row_count,
        selected_sheet,
        csv_output_path,
    )
    return csv_output_path


def copy_file_atomically(src_path, dst_path):
    """Copies src_path to dst_path atomically."""
    dst_dir = os.path.dirname(os.path.abspath(dst_path))
    os.makedirs(dst_dir, exist_ok=True)
    with tempfile.NamedTemporaryFile(
        "wb", dir=dst_dir, delete=False, suffix=".tmp"
    ) as tmp:
        with open(src_path, "rb") as fsrc:
            shutil.copyfileobj(fsrc, tmp)
        tmp_path = tmp.name
    os.replace(tmp_path, dst_path)
    logging.info("Copied %s to %s", src_path, dst_path)


def download_poverty_dataset(
    source_url=EDA_PPC_XLSX_URL,
    mirror_url=EDA_PPC_MIRROR_URL,
    input_file=None,
    output_xlsx_path=DEFAULT_OUTPUT_XLSX,
    output_csv_path=DEFAULT_INPUT_CSV,
    raw_csv_path=DEFAULT_RAW_CSV_FILE,
    max_retries=3,
    timeout=60,
):
    """Main workflow to download official EDA PPC workbook or ingest input file."""
    session = requests.Session()

    # Case 1: Local input file specified
    if input_file:
        if not os.path.exists(input_file) or os.path.getsize(input_file) == 0:
            raise FileNotFoundError(f"Input file not found or empty: {input_file}")
        logging.info("Using provided local input file: %s", input_file)
        if input_file.lower().endswith((".xlsx", ".xls")):
            copy_file_atomically(input_file, output_xlsx_path)
            extract_sheet_to_csv(input_file, output_csv_path)
        else:
            copy_file_atomically(input_file, output_csv_path)
        if raw_csv_path:
            copy_file_atomically(output_csv_path, raw_csv_path)
        return output_csv_path

    # Case 2: Download from web with primary and mirror fallback
    content = None
    urls_to_try = []
    if source_url:
        urls_to_try.append(source_url)
    if mirror_url and mirror_url != source_url:
        urls_to_try.append(mirror_url)

    last_err = None
    for url in urls_to_try:
        try:
            content = download_file(
                download_url=url,
                output_path=output_xlsx_path,
                session=session,
                max_retries=max_retries,
                timeout=timeout,
            )
            logging.info("Successfully downloaded workbook from %s", url)
            break
        except Exception as e:
            last_err = e
            logging.warning("Failed to download from %s: %s", url, e)

    # Fallback to existing input files if available
    if not content:
        for fallback in [output_xlsx_path, output_csv_path, os.path.join(MODULE_DIR, "test_data", "Poverty_input.csv")]:
            if os.path.exists(fallback) and os.path.getsize(fallback) > 0:
                logging.info("Falling back to existing local file: %s", fallback)
                if fallback.lower().endswith((".xlsx", ".xls")):
                    extract_sheet_to_csv(fallback, output_csv_path)
                elif fallback != output_csv_path:
                    copy_file_atomically(fallback, output_csv_path)
                if raw_csv_path and output_csv_path != raw_csv_path:
                    copy_file_atomically(output_csv_path, raw_csv_path)
                return output_csv_path

        raise RuntimeError(
            f"Failed to acquire dataset from all URLs: {urls_to_try}. Last error: {last_err}"
        ) from last_err

    # Extract Underlying_Data sheet to output_csv_path
    extract_sheet_to_csv(content, output_csv_path)
    if raw_csv_path:
        copy_file_atomically(output_csv_path, raw_csv_path)

    return output_csv_path


def main(argv):
    del argv  # Unused
    download_poverty_dataset(
        source_url=FLAGS.source_url,
        mirror_url=FLAGS.mirror_url,
        input_file=FLAGS.input_file,
        output_xlsx_path=FLAGS.output_xlsx_path,
        output_csv_path=FLAGS.output_csv_path,
        raw_csv_path=FLAGS.raw_csv_path,
        max_retries=FLAGS.max_retries,
        timeout=FLAGS.timeout,
    )


if __name__ == "__main__":
    app.run(main)
