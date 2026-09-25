# Copyright 2022 Google LLC
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#      http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.
"""
This Python script processes the CDC PRAMS consolidated MCH Indicators dataset
(Excel workbook), cleans it, and generates cleaned CSV, MCF, and TMCF files.
"""
import glob
import os
import sys
from copy import deepcopy
from absl import app, flags, logging
import numpy as np
import openpyxl
import pandas as pd

_CODEDIR = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(1, _CODEDIR)
sys.path.insert(1, os.path.join(_CODEDIR, '../../util/'))
from statvar_dcid_generator import get_statvar_dcid
from state_division_to_dcid import _PLACE_MAP
from statvar import statvar_col
from constants import (_MCF_TEMPLATE, _TMCF_TEMPLATE, DEFAULT_SV_PROP, _PROP,
                       _TIME, _INSURANCE, _CIGARETTES, PV_PROP, _YEAR)

_FLAGS = flags.FLAGS
default_input_path = os.path.join(_CODEDIR, "input_files")

flags.DEFINE_string("input_path", default_input_path,
                    "Path to input Excel file or directory containing files")
flags.DEFINE_string("output_path", None,
                    "Directory path where output files need to be written")
flags.DEFINE_list(
    "input_years", None,
    "Optional list of years to process (e.g. 2016,2017). "
    "Defaults to all available numeric sheets in the workbook.")

# Canonical 42 base Statistical Variable names in workbook column order
UNIQUE_STATVARS = list(dict.fromkeys(statvar_col.values()))


def _get_geo_map() -> dict:
    """Builds geographic mapping dictionary from site name to DCID."""
    geo_map = dict(_PLACE_MAP)
    geo_map.update({
        'Sites aggregated*': 'country/USA',
        'All Sites': 'country/USA',
        'New York City': 'geoId/3651000',
        'New York State': 'geoId/36',
        'Puerto Rico': 'geoId/72',
        'Northern Mariana Islands': 'geoId/69'
    })
    return geo_map


def prams(input_files: list, years: list = None) -> pd.DataFrame:
    """
    Parses CDC PRAMS Excel workbook(s) and produces observations DataFrame.

    Args:
        input_files (list): Paths to input files.
        years (list): Observation years to include.

    Returns:
        pd.DataFrame: Cleaned observations DataFrame.
    """
    geo_map = _get_geo_map()
    records = []

    excel_files = [f for f in input_files if f.endswith(('.xlsx', '.xls'))]
    if not excel_files:
        raise ValueError(f"No Excel files found in input files: {input_files}")

    for file_path in excel_files:
        logging.info("Reading Excel workbook: %s", file_path)
        wb = openpyxl.load_workbook(file_path, data_only=True)
        available_sheets = [s for s in wb.sheetnames if s.isdigit()]

        if years:
            target_sheets = [
                s for s in available_sheets if s in [str(y) for y in years]
            ]
        else:
            target_sheets = available_sheets

        for year in target_sheets:
            ws = wb[year]
            logging.info("Processing sheet year %s (%d rows)", year, ws.max_row)

            for r in range(6, ws.max_row + 1):
                site_raw = ws.cell(r, 1).value
                if not site_raw:
                    continue
                site = str(site_raw).strip()
                geo = geo_map.get(site)
                if not geo:
                    continue

                for ind_idx in range(len(UNIQUE_STATVARS)):
                    base_sv = UNIQUE_STATVARS[ind_idx]
                    base_col = 2 + ind_idx * 5

                    # 1. Sample Size (col + 0)
                    val_ss = ws.cell(r, base_col).value
                    if val_ss is not None and str(val_ss).strip() != '':
                        try:
                            val_ss_str = str(int(round(float(val_ss))))
                            records.append({
                                'Geo': geo,
                                'SV': f'SampleSize_Count{base_sv}',
                                'Year': str(year),
                                'Observation': val_ss_str,
                                'ScalingFactor': np.nan
                            })
                        except (ValueError, TypeError):
                            pass

                    # 2. Weighted Percent (col + 2)
                    val_pct = ws.cell(r, base_col + 2).value
                    if val_pct is not None and str(val_pct).strip() != '':
                        try:
                            val_pct_str = str(float(val_pct))
                            records.append({
                                'Geo': geo,
                                'SV': f'Percent{base_sv}',
                                'Year': str(year),
                                'Observation': val_pct_str,
                                'ScalingFactor': 100.0
                            })
                        except (ValueError, TypeError):
                            pass

                    # 3. Lower 95% Confidence Interval (col + 3)
                    val_lower = ws.cell(r, base_col + 3).value
                    if val_lower is not None and str(val_lower).strip() != '':
                        try:
                            val_lower_str = str(float(val_lower))
                            records.append({
                                'Geo': geo,
                                'SV': f'ConfidenceIntervalLowerLimit_Count{base_sv}',
                                'Year': str(year),
                                'Observation': val_lower_str,
                                'ScalingFactor': 100.0
                            })
                        except (ValueError, TypeError):
                            pass

                    # 4. Upper 95% Confidence Interval (col + 4)
                    val_upper = ws.cell(r, base_col + 4).value
                    if val_upper is not None and str(val_upper).strip() != '':
                        try:
                            val_upper_str = str(float(val_upper))
                            records.append({
                                'Geo': geo,
                                'SV': f'ConfidenceIntervalUpperLimit_Count{base_sv}',
                                'Year': str(year),
                                'Observation': val_upper_str,
                                'ScalingFactor': 100.0
                            })
                        except (ValueError, TypeError):
                            pass

    df = pd.DataFrame(records)
    if df.empty:
        raise ValueError(
            "No observation records were extracted from Excel workbook(s)")

    return df


class USPrams:
    """
    Class to process CDC PRAMS data, generate schema MCF, template MCF,
    and cleaned CSV output files.
    """

    def __init__(self,
                 input_files: list,
                 output_location: str = 'output',
                 output_file_name: str = 'PRAMS',
                 years: list = None):
        self.input_files = input_files
        self.output_location = output_location
        self.years = years
        self.cleaned_csv_file_path = os.path.join(self.output_location,
                                                  f"{output_file_name}.csv")
        self.mcf_file_path = os.path.join(self.output_location,
                                          f"{output_file_name}.mcf")
        self.tmcf_file_path = os.path.join(self.output_location,
                                           f"{output_file_name}.tmcf")

    def _generate_tmcf(self) -> None:
        """Generates Template MCF file atomically."""
        tmp_tmcf_file = self.tmcf_file_path + ".tmp"
        with open(tmp_tmcf_file, 'w', encoding='utf-8') as f_out:
            f_out.write(_TMCF_TEMPLATE.rstrip('\n') + '\n')
        os.replace(tmp_tmcf_file, self.tmcf_file_path)

    def _generate_mcf(self, sv_names: list, mcf_file_path: str) -> dict:
        """
        Generates schema MCF file and returns mapping from dummy SV to DCID.
        """
        mcf_nodes = []
        dcid_nodes = {}
        for sv in sv_names:
            pvs = []
            dcid = sv
            sv_prop = [prop.strip() for prop in sv.split(" ")]
            sv_pvs = deepcopy(DEFAULT_SV_PROP)

            for prop in sv_prop:
                statVar = insurance = time = cigarettes = prop_val = prop
                for old, new in _PROP.items():
                    statVar = statVar.replace(old, new)
                for old, new in _INSURANCE.items():
                    insurance = insurance.replace(old, new)
                for old, new in _TIME.items():
                    time = time.replace(old, new)
                for old, new in _CIGARETTES.items():
                    cigarettes = cigarettes.replace(old, new)
                for old, new in PV_PROP.items():
                    prop_val = prop_val.replace(old, new)

                if "SampleSize" in prop:
                    sv_pvs["measuredProperty"] = "dcs:count"
                    sv_pvs["statType"] = "dcs:sampleSize"
                    pvs.append("measuredProperty: dcs:count")
                    pvs.append("statType: dcs:sampleSize")

                if "Percent" in prop:
                    sv_pvs["measuredProperty"] = "dcs:percent"
                    sv_pvs["statType"] = "dcs:measuredValue"
                    sv_pvs["measurementDenominator"] = (
                        "dcs:Count_BirthEvent_LiveBirth"
                    )
                    pvs.append("measuredProperty: dcs:count")
                    pvs.append("statType: dcs:measuredValue")
                    pvs.append(
                        "measurementDenominator: dcs:Count_BirthEvent_LiveBirth"
                    )

                if "ConfidenceIntervalLowerLimit" in prop:
                    sv_pvs["measuredProperty"] = "dcs:percent"
                    sv_pvs["statType"] = "dcs:confidenceIntervalLowerLimit"
                    sv_pvs["measurementDenominator"] = (
                        "dcs:Count_BirthEvent_LiveBirth"
                    )
                    pvs.append("measuredProperty: dcs:count")
                    pvs.append("statType: dcs:confidenceIntervalLowerLimit")
                    pvs.append(
                        "measurementDenominator: dcs:Count_BirthEvent_LiveBirth"
                    )

                if "ConfidenceIntervalUpperLimit" in prop:
                    sv_pvs["measuredProperty"] = "dcs:percent"
                    sv_pvs["statType"] = "dcs:confidenceIntervalUpperLimit"
                    sv_pvs["measurementDenominator"] = (
                        "dcs:Count_BirthEvent_LiveBirth"
                    )
                    pvs.append("measuredProperty: dcs:count")
                    pvs.append("statType: dcs:confidenceIntervalUpperLimit")
                    pvs.append(
                        "measurementDenominator: dcs:Count_BirthEvent_LiveBirth"
                    )

                if "MultivitaminUseMoreThan4TimesAWeek" in prop:
                    prop = prop[0].lower() + prop[1:]
                    sv_pvs["mothersHealthPrevention"] = f"dcs:{statVar}"
                    sv_pvs["healthPreventionActionFrequency"] = f"dcs:{time}"
                    pvs.append(f"mothersHealthPrevention: dcs:{statVar}")
                    pvs.append(f"healthPreventionActionFrequency: dcs:{time}")

                if "Underweight" in prop or "Overweight" in prop or\
                    "Obese" in prop:
                    sv_pvs["mothersHealthBehavior"] = f"dcs:{statVar}"
                    pvs.append(f"mothersHealthBehavior: dcs:{statVar}")

                if "HealthCareVisit12MonthsBeforePregnancy" in prop or\
                        "PrenatalCareInFirstTrimester" in prop or\
                        "FluShot12MonthsBeforeDelivery" in prop or\
                        "MaternalCheckupPostpartum" in prop or\
                        "TeethCleanedByDentistOrHygienist" in prop:
                    prop = prop[0].lower() + prop[1:]
                    sv_pvs["mothersHealthPrevention"] = f"dcs:{statVar}"
                    sv_pvs["timePeriodRelativeToPregnancy"] = f"dcs:{time}"
                    pvs.append(f"mothersHealthPrevention: dcs:{statVar}")
                    pvs.append(f"timePeriodRelativeToPregnancy: dcs:{time}")

                elif "CigaretteSmoking3MonthsBeforePregnancy" in prop or\
                    "CigaretteSmokingLast3MonthsOfPregnancy" in prop or \
                    "CigaretteSmokingPostpartum" in prop or\
                    "ECigaretteSmoking3MonthsBeforePregnancy" in prop or\
                    "ECigaretteSmokingLast3MonthsOfPregnancy" in prop:
                    sv_pvs["tobaccoUsageType"] = f"dcs:{cigarettes}"
                    sv_pvs["mothersHealthBehavior"] = f"dcs:{statVar}"
                    sv_pvs["timePeriodRelativeToPregnancy"] = f"{time}"
                    pvs.append(f"tobaccoUsageType: dcs:{cigarettes}")
                    pvs.append(f"mothersHealthBehavior: dcs:{statVar}")
                    pvs.append(f"timePeriodRelativeToPregnancy: dcs:{time}")

                elif "HookahInLast2Years" in prop or\
                     "HeavyDrinking3MonthsBeforePregnancy" in prop:
                    sv_pvs["mothersHealthBehavior"] = f"dcs:{statVar}"
                    sv_pvs["timePeriodRelativeToPregnancy"] = f"{time}"
                    pvs.append(f"mothersHealthBehavior: dcs:{statVar}")
                    pvs.append(f"timePeriodRelativeToPregnancy: dcs:{time}")

                elif "IntimatePartnerViolenceByCurrentOrExPartnerOr"+\
                    "CurrentOrExHusband" in prop:
                    sv_pvs["intimatePartnerViolence"] = f"dcs:{statVar}"
                    sv_pvs["timePeriodRelativeToPregnancy"] = f"{time}"
                    pvs.append(f"intimatePartnerViolence: dcs:{statVar}")
                    pvs.append(f"timePeriodRelativeToPregnancy: dcs:{time}")

                elif "MistimedPregnancy" in prop or\
                    "UnwantedPregnancy" in prop or\
                    "UnsureIfWantedPregnancy" in prop or\
                     "IntendedPregnancy" in prop:
                    sv_pvs["pregnancyIntention"] = f"dcs:{cigarettes}"
                    pvs.append(f"pregnancyIntention: dcs:{cigarettes}")

                elif "AnyPostpartumFamilyPlanning" in prop or\
                    "MaleOrFemaleSterilization" in prop or\
                    "LongActingReversibleContraceptiveMethods" in prop or\
                    "ModeratelyEffectiveContraceptiveMethods" in prop or\
                    "LeastEffectiveContraceptiveMethods" in prop:
                    sv_pvs["postpartumFamilyPlanning"] = f"dcs:{cigarettes}"
                    pvs.append(f"postpartumFamilyPlanning: dcs:{cigarettes}")

                elif "CDC_SelfReportedDepression3MonthsBeforePregnancy" in prop\
                     or "CDC_SelfReportedDepressionDuringPregnancy" in prop or\
                    "CDC_SelfReportedDepressionPostpartum" in prop:
                    sv_pvs["mothersHealthCondition"] = f"dcs:{statVar}"
                    sv_pvs["timePeriodRelativeToPregnancy"] = f"dcs:{time}"
                    pvs.append(f"mothersHealthCondition: dcs:{statVar}")
                    pvs.append(f"timePeriodRelativeToPregnancy: dcs:{time}")

                elif "healthInsuranceStatusOneMonthBeforePregnancy"+\
                    "PrivateInsurance" in prop or\
                    "healthInsuranceStatusOneMonthBeforePregnancy"+\
                        "Medicaid" in prop or\
                    "healthInsuranceStatusOneMonthBeforePregnancy"+\
                        "NoInsurance" in prop:
                    sv_pvs["healthInsuranceStatusOneMonthBeforePregnancy"] = (
                        f"dcs:{statVar}")
                    sv_pvs["timePeriodRelativeToPregnancy"] = f"dcs:{time}"
                    pvs.append(
                        f"healthInsuranceStatusOneMonthBeforePregnancy: dcs:{statVar}"
                    )
                    pvs.append(f"timePeriodRelativeToPregnancy: dcs:{time}")

                elif "healthInsuranceStatusForPrenatalCare"+\
                    "PrivateInsurance" in prop or\
                    "healthInsuranceStatusForPrenatalCareMedicaid" in prop or\
                    "healthInsuranceStatusForPrenatalCareNoInsurance" in prop:
                    sv_pvs["healthInsuranceStatusForPrenatalCare"] = (
                        f"dcs:{statVar}")
                    pvs.append(
                        f"healthInsuranceStatusForPrenatalCare: dcs:{statVar}")

                elif "healthInsuranceStatusPostpartumPrivateInsurance" in prop\
                    or "healthInsuranceStatusPostpartumMedicaid" in prop or\
                    "healthInsuranceStatusPostpartumNoInsurance" in prop:
                    sv_pvs["healthInsuranceStatusPostpartum"] = f"dcs:{insurance}"
                    sv_pvs["timePeriodRelativeToPregnancy"] = f"dcs:{time}"
                    pvs.append(
                        f"healthInsuranceStatusPostpartum: dcs:{insurance}")
                    pvs.append(f"timePeriodRelativeToPregnancy: dcs:{time}")

                elif "BabyMostOftenLaidOnBackToSleep" in prop:
                    sv_pvs["infantSleepPractice"] = f"dcs:{cigarettes}"
                    pvs.append(f"infantSleepPractice: dcs:{cigarettes}")

                elif "EverBreastfed" in prop:
                    sv_pvs["breastFeedingPractice"] = f"dcs:{cigarettes}"
                    pvs.append(f"breastFeedingPractice: dcs:{cigarettes}")

                elif "AnyBreastfeedingAt8Weeks" in prop:
                    sv_pvs["breastFeedingPractice"] = f"dcs:{cigarettes}"
                    sv_pvs["timePeriodRelativeToPregnancy"] = f"dcs:{time}"
                    pvs.append(f"breastFeedingPractice: dcs:{cigarettes}")
                    pvs.append(f"timePeriodRelativeToPregnancy: dcs:{time}")

            resolved_dcid = get_statvar_dcid(sv_pvs)
            dcid_nodes[dcid] = resolved_dcid
            mcf_nodes.append(
                _MCF_TEMPLATE.format(dcid=resolved_dcid,
                                     xtra_pvs='\n'.join(pvs)))
        mcf = '\n'.join(mcf_nodes)

        tmp_mcf_file = mcf_file_path + ".tmp"
        with open(tmp_mcf_file, 'w+', encoding='utf-8') as f_out:
            f_out.write(mcf.rstrip('\n') + '\n')
        os.replace(tmp_mcf_file, mcf_file_path)
        return dcid_nodes

    def process(self) -> None:
        """
        Executes pipeline: extracts Excel data, resolves StatVars,
        and atomically outputs CSV, MCF, and TMCF files.
        """
        df = prams(self.input_files, self.years)
        sv_names = sorted(df["SV"].unique().tolist())

        output_path = os.path.dirname(self.cleaned_csv_file_path)
        if output_path and not os.path.exists(output_path):
            os.makedirs(output_path, exist_ok=True)

        updated_sv = self._generate_mcf(sv_names, self.mcf_file_path)
        df["SV"] = df["SV"].map(updated_sv)
        if df["SV"].isna().any():
            unmapped = df[df["SV"].isna()]["SV"].unique().tolist()
            logging.error("Unmapped Statistical Variables detected: %s",
                          unmapped)
            raise ValueError(
                f"Unmapped Statistical Variables detected: {unmapped}")

        self._generate_tmcf()
        df["Observation"] = df["Observation"].replace(to_replace={'': pd.NA})
        df = df.dropna(subset=['Observation'])
        df = df.drop_duplicates(subset=['Geo', 'SV', 'Year'], keep='last')
        df = df.sort_values(by=['Geo', 'SV', 'Year'])

        tmp_csv_file = self.cleaned_csv_file_path + ".tmp"
        df.to_csv(tmp_csv_file, index=False)
        os.replace(tmp_csv_file, self.cleaned_csv_file_path)
        logging.info("Successfully produced: %s (%d records)",
                     self.cleaned_csv_file_path, len(df))


def main(_):
    input_path = _FLAGS.input_path
    if not os.path.exists(input_path):
        logging.fatal("Input path not found: %s", input_path)

    if os.path.isfile(input_path):
        ip_files = [input_path]
    else:
        ip_files = sorted(glob.glob(os.path.join(input_path, "*.xlsx")))
        if not ip_files:
            ip_files = sorted(glob.glob(os.path.join(input_path, "*.xls")))

    if not ip_files:
        logging.fatal("No Excel input files found in: %s", input_path)

    output_dir = _FLAGS.output_path or os.path.join(_CODEDIR, "output")
    os.makedirs(output_dir, exist_ok=True)

    loader = USPrams(ip_files,
                     output_location=output_dir,
                     years=_FLAGS.input_years)
    loader.process()


if __name__ == '__main__':
    app.run(main)
